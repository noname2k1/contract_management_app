import os

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from utils import (
    auto_fit_row_heights,
    format_money,
    parse_number,
    prepare_product_area,
    restore_template_drawing,
)

MAX_VENDORS = 3


# ============================================================
# XUẤT EXCEL
# ============================================================
def export_duyet_gia(
    data,
    template_path,
    output_path,
):
    if not os.path.exists(template_path):
        raise FileNotFoundError(
            "Không tìm thấy file mẫu:\n" + os.path.abspath(template_path)
        )

    wb = load_workbook(template_path)

    if "Bộ 01" not in wb.sheetnames:
        raise ValueError("File Excel không có sheet 'Bộ 01'.")

    ws = wb["Bộ 01"]

    vendors = data["vendors"]
    products = data["products"]

    if not vendors:
        raise ValueError("Phải có ít nhất 1 nhà cung cấp.")

    if len(vendors) > MAX_VENDORS:
        raise ValueError(f"Mẫu này chỉ hỗ trợ tối đa {MAX_VENDORS} nhà cung cấp.")

    if not products:
        raise ValueError("Phải có ít nhất 1 hạng mục.")

    # ========================================================
    # THÔNG TIN CHUNG
    # ========================================================

    so_duyet = data["so_duyet"].strip()
    ngay = data["ngay"].strip()
    thang = data["thang"].strip()
    nam = data["nam"].strip()

    ten_duyet_gia = data["ten_duyet_gia"].strip()

    muc_dich = data["muc_dich_mua_sam"].strip()

    dia_diem = data["dia_diem"].strip()

    ws["A4"] = f"Số: {so_duyet}/KTCN"

    ws["F4"] = f"{dia_diem}, ngày     tháng      năm     "

    ws["A6"] = "Ngày     tháng      năm     "

    ws["F5"] = f"ĐÁNH GIÁ VÀ ĐỀ XUẤT\nLựa chọn nhà cung cấp mua sắm\n{ten_duyet_gia}"

    # ========================================================
    # DANH SÁCH NHÀ CUNG CẤP
    # ========================================================

    vendor_sentence = (
        "        Căn cứ yêu cầu mua sắm "
        f"thiết bị, dịch vụ {muc_dich}, "
        "Phòng KTCN đã yêu cầu báo giá "
        "cung cấp của "
        f"{len(vendors):02d} nhà cung cấp sau:"
    )

    for i, vendor in enumerate(
        vendors,
        start=1,
    ):
        vendor_sentence += f"\n{i}. {vendor['name'].strip()}"

    ws["A13"] = vendor_sentence

    # ========================================================
    # TÊN NCC
    # ========================================================

    for col in range(5, 8):
        cell = ws.cell(
            16,
            col,
        )

        if isinstance(
            cell,
            MergedCell,
        ):
            raise ValueError(
                f"Ô {cell.coordinate} đang là MergedCell. Kiểm tra merge ở hàng 16."
            )

        cell.value = None

    for index, vendor in enumerate(vendors):
        col = 5 + index

        ws.cell(
            16,
            col,
        ).value = vendor["name"].strip()

    # ========================================================
    # KHU VỰC SẢN PHẨM
    # ========================================================

    first_product_row = 18

    prepare_product_area(
        ws,
        len(products),
    )

    evaluation_start = first_product_row + len(products)

    for index, product in enumerate(products):
        row = first_product_row + index

        ws.cell(
            row,
            1,
        ).value = f"1.{index + 1}"

        ws.cell(
            row,
            2,
        ).value = product["name"].strip()

        quantity = product["quantity"].strip()

        try:
            quantity_value = float(quantity)

            if quantity_value.is_integer():
                quantity_value = int(quantity_value)

        except ValueError:
            quantity_value = quantity

        ws.cell(
            row,
            3,
        ).value = quantity_value

        ws.cell(
            row,
            4,
        ).value = product["unit"].strip()

        prices = []

        for vendor_index in range(MAX_VENDORS):
            col = 5 + vendor_index

            if vendor_index < len(vendors):
                price = parse_number(product["prices"][vendor_index])

                ws.cell(
                    row,
                    col,
                ).value = price

                ws.cell(
                    row,
                    col,
                ).number_format = "#,##0"

                prices.append(price)

            else:
                ws.cell(
                    row,
                    col,
                ).value = None

        if not prices:
            raise ValueError(f"Hạng mục {index + 1} chưa có giá.")

        min_price = min(prices)

        ws.cell(
            row,
            8,
        ).value = min_price

        ws.cell(
            row,
            8,
        ).number_format = "#,##0"

        min_index = prices.index(min_price)

        ws.cell(
            row,
            9,
        ).value = vendors[min_index]["name"].strip()

        ws.cell(
            row,
            10,
        ).value = product.get(
            "note",
            "",
        ).strip()

    # ========================================================
    # ĐÁNH GIÁ NCC
    # ========================================================

    row_capacity = evaluation_start

    row_technical = evaluation_start + 1

    row_delivery = evaluation_start + 2

    row_time = evaluation_start + 3

    row_payment = evaluation_start + 4

    row_proposal = evaluation_start + 5

    labels = {
        row_capacity: "Đánh giá về năng lực và KN",
        row_technical: "Đánh giá về đáp ứng yêu cầu KT",
        row_delivery: "Địa điểm giao hàng/thực hiện",
        row_time: "Thời gian thực hiện hợp đồng",
        row_payment: "Thời hạn thanh toán",
    }

    for row, label in labels.items():
        ws.cell(
            row,
            1,
        ).value = row - evaluation_start + 2

        ws.cell(
            row,
            2,
        ).value = label

    for vendor_index, vendor in enumerate(vendors):
        col = 5 + vendor_index

        ws.cell(
            row_capacity,
            col,
        ).value = vendor["nang_luc"].strip()

        ws.cell(
            row_technical,
            col,
        ).value = vendor["ky_thuat"].strip()

        ws.cell(
            row_delivery,
            col,
        ).value = vendor["dia_diem"].strip()

        ws.cell(
            row_time,
            col,
        ).value = vendor["thoi_gian"].strip()

        ws.cell(
            row_payment,
            col,
        ).value = vendor["cach_thanh_toan"].strip()

    for vendor_index in range(
        len(vendors),
        MAX_VENDORS,
    ):
        col = 5 + vendor_index

        for row in labels:
            ws.cell(
                row,
                col,
            ).value = None

    # ========================================================
    # TỔNG GIÁ
    # ========================================================

    vendor_totals = [0 for _ in vendors]

    for product in products:
        quantity = parse_number(product["quantity"])

        for i in range(len(vendors)):
            vendor_totals[i] += parse_number(product["prices"][i]) * quantity

    selected_index = vendor_totals.index(min(vendor_totals))

    selected_vendor = vendors[selected_index]["name"].strip()

    selected_total = vendor_totals[selected_index]

    # ========================================================
    # ĐỀ XUẤT
    # ========================================================

    proposal = (
        "        Phòng KTCN đề xuất "
        "Thủ trưởng Nhà máy cho mua "
        "các hạng mục từ "
        f"{selected_vendor} vì có tổng "
        "giá chào thấp nhất trong các "
        "nhà cung cấp được đánh giá, "
        "đáp ứng các yêu cầu về vật tư, "
        "kỹ thuật và khả năng thực hiện."
    )

    target_merge = f"A{row_proposal}:J{row_proposal}"

    already_merged = any(str(rng) == target_merge for rng in ws.merged_cells.ranges)

    if not already_merged:
        for rng in list(ws.merged_cells.ranges):
            min_col, min_row, max_col, max_row = rng.bounds

            if (
                min_row == row_proposal
                and max_row == row_proposal
                and min_col <= 1 <= max_col
            ):
                ws.unmerge_cells(str(rng))

        ws.merge_cells(target_merge)

    ws.cell(
        row_proposal,
        1,
    ).value = proposal

    # ========================================================
    # COMMENT TỔNG GIÁ
    # ========================================================

    from openpyxl.comments import Comment

    for i, total in enumerate(vendor_totals):
        col = 5 + i

        ws.cell(
            16,
            col,
        ).comment = Comment(
            f"Tổng giá chưa VAT: {format_money(total)} đồng",
            "Python",
        )

    # ========================================================
    # FORMAT
    # ========================================================

    auto_fit_row_heights(
        ws,
        min_row=17,
        max_row=ws.max_row,
        min_height=15.0,
        max_height=409.0,
    )

    ws.sheet_view.showGridLines = False

    os.makedirs(
        os.path.dirname(os.path.abspath(output_path)),
        exist_ok=True,
    )

    wb.save(output_path)

    restore_template_drawing(
        template_path,
        output_path,
    )

    return {
        "output_path": os.path.abspath(output_path),
        "selected_vendor": selected_vendor,
        "selected_total": selected_total,
        "vendor_totals": vendor_totals,
    }
