import os
import re
import tkinter as tk
from copy import deepcopy
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from tkinterdnd2 import DND_FILES

from zoneinfo import ZoneInfo

VN_TZ = ZoneInfo("Asia/Ho_Chi_Minh")

from utils import (
    auto_fit_row_heights,
    format_money,
    number_to_vietnamese_words,
    open_file,
    parse_number,
    prepare_product_area,
    restore_template_drawing,
)
from utils.document_utils import generate_docx

# ============================================================
# CẤU HÌNH
# ============================================================

DEFAULT_OUTPUT_DIR = "./outputs"
TEMPLATE_FILE = "./templates/duyet_gia.xlsx"
DEFAULT_SERVICE_TEMPLATE = "./templates/DV.docx"
DEFAULT_SALE_TEMPLATE = "./templates/TB.docx"
DEFAULT_CONTRACT_OUTPUT_DIR = "./outputs/contracts"

MAX_VENDORS = 3

# ============================================================
# XUẤT EXCEL
# ============================================================


def export_duyet_gia(
    data,
    template_path,
    output_path,
):
    if not os.path.exists(template_path):
        raise FileNotFoundError(
            "Không tìm thấy file mẫu:\n" + os.path.abspath(template_path)
        )

    wb = load_workbook(template_path)

    if "Bộ 01" not in wb.sheetnames:
        raise ValueError("File Excel không có sheet 'Bộ 01'.")

    ws = wb["Bộ 01"]

    vendors = data["vendors"]
    products = data["products"]

    if not vendors:
        raise ValueError("Phải có ít nhất 1 nhà cung cấp.")

    if len(vendors) > MAX_VENDORS:
        raise ValueError(f"Mẫu này chỉ hỗ trợ tối đa {MAX_VENDORS} nhà cung cấp.")

    if not products:
        raise ValueError("Phải có ít nhất 1 hạng mục.")

    # ========================================================
    # THÔNG TIN CHUNG
    # ========================================================

    so_duyet = data["so_duyet"].strip()
    ngay = data["ngay"].strip()
    thang = data["thang"].strip()
    nam = data["nam"].strip()

    ten_duyet_gia = data["ten_duyet_gia"].strip()

    muc_dich = data["muc_dich_mua_sam"].strip()

    dia_diem = data["dia_diem"].strip()

    ws["A4"] = f"Số: {so_duyet}/KTCN"

    ws["F4"] = f"{dia_diem}, ngày {ngay} tháng {thang} năm {nam}"

    ws["A6"] = f"Ngày {ngay} tháng {thang} năm {nam}"

    ws["F5"] = f"ĐÁNH GIÁ VÀ ĐỀ XUẤT\nLựa chọn nhà cung cấp mua sắm\n{ten_duyet_gia}"

    # ========================================================
    # DANH SÁCH NHÀ CUNG CẤP
    # ========================================================

    vendor_sentence = (
        "        Căn cứ yêu cầu mua sắm "
        f"thiết bị, dịch vụ {muc_dich}, "
        "Phòng KTCN đã yêu cầu báo giá "
        "cung cấp của "
        f"{len(vendors):02d} nhà cung cấp sau:"
    )

    for i, vendor in enumerate(
        vendors,
        start=1,
    ):
        vendor_sentence += f"\n{i}. {vendor['name'].strip()}"

    ws["A13"] = vendor_sentence

    # ========================================================
    # TÊN NCC
    # ========================================================

    for col in range(5, 8):
        cell = ws.cell(
            16,
            col,
        )

        if isinstance(
            cell,
            MergedCell,
        ):
            raise ValueError(
                f"Ô {cell.coordinate} đang là MergedCell. Kiểm tra merge ở hàng 16."
            )

        cell.value = None

    for index, vendor in enumerate(vendors):
        col = 5 + index

        ws.cell(
            16,
            col,
        ).value = vendor["name"].strip()

    # ========================================================
    # KHU VỰC SẢN PHẨM
    # ========================================================

    first_product_row = 18

    prepare_product_area(
        ws,
        len(products),
    )

    evaluation_start = first_product_row + len(products)

    for index, product in enumerate(products):
        row = first_product_row + index

        ws.cell(
            row,
            1,
        ).value = f"1.{index + 1}"

        ws.cell(
            row,
            2,
        ).value = product["name"].strip()

        quantity = product["quantity"].strip()

        try:
            quantity_value = float(quantity)

            if quantity_value.is_integer():
                quantity_value = int(quantity_value)

        except ValueError:
            quantity_value = quantity

        ws.cell(
            row,
            3,
        ).value = quantity_value

        ws.cell(
            row,
            4,
        ).value = product["unit"].strip()

        prices = []

        for vendor_index in range(MAX_VENDORS):
            col = 5 + vendor_index

            if vendor_index < len(vendors):
                price = parse_number(product["prices"][vendor_index])

                ws.cell(
                    row,
                    col,
                ).value = price

                ws.cell(
                    row,
                    col,
                ).number_format = "#,##0"

                prices.append(price)

            else:
                ws.cell(
                    row,
                    col,
                ).value = None

        if not prices:
            raise ValueError(f"Hạng mục {index + 1} chưa có giá.")

        min_price = min(prices)

        ws.cell(
            row,
            8,
        ).value = min_price

        ws.cell(
            row,
            8,
        ).number_format = "#,##0"

        min_index = prices.index(min_price)

        ws.cell(
            row,
            9,
        ).value = vendors[min_index]["name"].strip()

        ws.cell(
            row,
            10,
        ).value = product.get(
            "note",
            "",
        ).strip()

    # ========================================================
    # ĐÁNH GIÁ NCC
    # ========================================================

    row_capacity = evaluation_start

    row_technical = evaluation_start + 1

    row_delivery = evaluation_start + 2

    row_time = evaluation_start + 3

    row_payment = evaluation_start + 4

    row_proposal = evaluation_start + 5

    labels = {
        row_capacity: "Đánh giá về năng lực và KN",
        row_technical: "Đánh giá về đáp ứng yêu cầu KT",
        row_delivery: "Địa điểm giao hàng/thực hiện",
        row_time: "Thời gian thực hiện hợp đồng",
        row_payment: "Thời hạn thanh toán",
    }

    for row, label in labels.items():
        ws.cell(
            row,
            1,
        ).value = row - evaluation_start + 2

        ws.cell(
            row,
            2,
        ).value = label

    for vendor_index, vendor in enumerate(vendors):
        col = 5 + vendor_index

        ws.cell(
            row_capacity,
            col,
        ).value = vendor["nang_luc"].strip()

        ws.cell(
            row_technical,
            col,
        ).value = vendor["ky_thuat"].strip()

        ws.cell(
            row_delivery,
            col,
        ).value = vendor["dia_diem"].strip()

        ws.cell(
            row_time,
            col,
        ).value = vendor["thoi_gian"].strip()

        ws.cell(
            row_payment,
            col,
        ).value = vendor["cach_thanh_toan"].strip()

    for vendor_index in range(
        len(vendors),
        MAX_VENDORS,
    ):
        col = 5 + vendor_index

        for row in labels:
            ws.cell(
                row,
                col,
            ).value = None

    # ========================================================
    # TỔNG GIÁ
    # ========================================================

    vendor_totals = [0 for _ in vendors]

    for product in products:
        quantity = parse_number(product["quantity"])

        for i in range(len(vendors)):
            vendor_totals[i] += parse_number(product["prices"][i]) * quantity

    selected_index = vendor_totals.index(min(vendor_totals))

    selected_vendor = vendors[selected_index]["name"].strip()

    selected_total = vendor_totals[selected_index]

    # ========================================================
    # ĐỀ XUẤT
    # ========================================================

    proposal = (
        "        Phòng KTCN đề xuất "
        "Thủ trưởng Nhà máy cho mua "
        "các hạng mục từ "
        f"{selected_vendor} vì có tổng "
        "giá chào thấp nhất trong các "
        "nhà cung cấp được đánh giá, "
        "đáp ứng các yêu cầu về vật tư, "
        "kỹ thuật và khả năng thực hiện."
    )

    target_merge = f"A{row_proposal}:J{row_proposal}"

    already_merged = any(str(rng) == target_merge for rng in ws.merged_cells.ranges)

    if not already_merged:
        for rng in list(ws.merged_cells.ranges):
            min_col, min_row, max_col, max_row = rng.bounds

            if (
                min_row == row_proposal
                and max_row == row_proposal
                and min_col <= 1 <= max_col
            ):
                ws.unmerge_cells(str(rng))

        ws.merge_cells(target_merge)

    ws.cell(
        row_proposal,
        1,
    ).value = proposal

    # ========================================================
    # COMMENT TỔNG GIÁ
    # ========================================================

    from openpyxl.comments import Comment

    for i, total in enumerate(vendor_totals):
        col = 5 + i

        ws.cell(
            16,
            col,
        ).comment = Comment(
            f"Tổng giá chưa VAT: {format_money(total)} đồng",
            "Python",
        )

    # ========================================================
    # FORMAT
    # ========================================================

    auto_fit_row_heights(
        ws,
        min_row=17,
        max_row=ws.max_row,
        min_height=15.0,
        max_height=409.0,
    )

    ws.sheet_view.showGridLines = False

    os.makedirs(
        os.path.dirname(os.path.abspath(output_path)),
        exist_ok=True,
    )

    wb.save(output_path)

    restore_template_drawing(
        template_path,
        output_path,
    )

    return {
        "output_path": os.path.abspath(output_path),
        "selected_vendor": selected_vendor,
        "selected_total": selected_total,
        "vendor_totals": vendor_totals,
    }


# ============================================================
# WORD
# ============================================================


def generate_contract_docx(self):
    try:
        # ============================================================
        # 1. LẤY TOÀN BỘ DỮ LIỆU TỪ GIAO DIỆN
        # ============================================================
        data = self.get_data()
        # ============================================================
        # 2. KIỂM TRA TEMPLATE WORD
        # ============================================================
        template_file = getattr(self, "template_contract_file", None)

        if not template_file:
            template_file = DEFAULT_SERVICE_TEMPLATE

        if not os.path.exists(template_file):
            raise FileNotFoundError(
                "Không tìm thấy file Word mẫu:\n\n" + str(template_file)
            )

        # ============================================================
        # 3. LẤY DANH SÁCH SẢN PHẨM
        # ============================================================

        raw_products = data.get("products", [])

        products = []

        for index, product in enumerate(raw_products, start=1):
            # --------------------------------------------------------
            # Tên sản phẩm
            # --------------------------------------------------------

            name = product.get("name", "")

            # --------------------------------------------------------
            # Đơn vị
            # --------------------------------------------------------

            unit = product.get("unit", "")

            # --------------------------------------------------------
            # Số lượng
            # --------------------------------------------------------

            quantity = parse_number(product.get("quantity", 0))

            # --------------------------------------------------------
            # Giá
            #
            # get_data() của bạn có thể trả:
            #
            # price
            #
            # hoặc:
            #
            # prices = [...]
            # --------------------------------------------------------

            price = product.get("price", None)

            if price is None:
                prices = product.get("prices", [])

                # Lấy giá của nhà cung cấp được chọn
                selected_index = data.get("selected_vendor_index", 0)

                try:
                    price = prices[selected_index]
                except Exception:  # noqa: BLE001
                    price = prices[0] if prices else 0

            price = parse_number(price)

            # --------------------------------------------------------
            # Thành tiền
            # --------------------------------------------------------

            amount = quantity * price

            # --------------------------------------------------------
            # STT
            # --------------------------------------------------------

            product_id = product.get("id", index)

            # --------------------------------------------------------
            # Đưa vào context
            # --------------------------------------------------------

            products.append(
                {
                    "id": product_id,
                    "name": name,
                    "unit": unit,
                    "quantity": quantity,
                    "price": format_money(price),
                    "amount": format_money(amount),
                }
            )

        # ============================================================
        # 4. TÍNH TỔNG
        # ============================================================

        sub_total = sum(parse_number(product["amount"]) for product in products)

        # VAT 8%
        vat = sub_total * 8 / 100

        grand_total = sub_total + vat

        # ============================================================
        # 5. LẤY CÁC BIẾN CHUNG
        # ============================================================

        context = {}

        # ------------------------------------------------------------
        # Copy toàn bộ data đơn giản
        # ------------------------------------------------------------

        for key, value in data.items():
            # Không đưa các object phức tạp
            # vào context nếu không cần.
            if key in (
                "products",
                "vendors",
            ):
                continue

            context[key] = value

        # ============================================================
        # 6. ĐƯA CONTRACT VARS VÀO CONTEXT
        # ============================================================

        contract_vars = data.get("contract_vars", {})

        if isinstance(contract_vars, dict):
            for key, value in contract_vars.items():
                context[key] = value

        # ============================================================
        # 7. CÁC BIẾN GIAO DIỆN QUAN TRỌNG
        # ============================================================

        context.update(
            {
                "products": products,
                "sub_total": format_money(sub_total),
                "VAT": format_money(vat),
                "grand_total": format_money(grand_total),
                "grand_total_text": number_to_vietnamese_words(grand_total),
            }
        )

        # ============================================================
        # 8. TẠO TÊN FILE
        # ============================================================

        ma_hd = context.get("ma_hd", data.get("ma_hd", ""))

        if not ma_hd:
            ma_hd = "Hop_dong"

        # Làm sạch tên file
        ma_hd = re.sub(r'[\\/:*?"<>|]+', "_", str(ma_hd))

        output_filename = f"HD{context.get('loai_hd', 'DV')}_{ma_hd}{datetime.now(VN_TZ).strftime('%Y%m%d_%H%M%S')}.docx"

        # ============================================================
        # 9. XUẤT WORD BẰNG generate_docx()
        # ============================================================

        output_path = generate_docx(
            template_file=template_file,
            context_data=context,
            output_filename=output_filename,
            output_dir=DEFAULT_CONTRACT_OUTPUT_DIR,
        )

        # ============================================================
        # 10. THÔNG BÁO
        # ============================================================
        # print(context)
        messagebox.showinfo(
            "Xuất hợp đồng thành công",
            f"Đã tạo:\n"
            f"{output_path}\n\n"
            f"Đối tác: "
            f"{context.get('ma_doi_tac')} - {context.get('ten_doi_tac')}\n"
            f"Giá trị trước VAT: "
            f"{format_money(sub_total)} đồng\n"
            f"VAT 8%: "
            f"{format_money(vat)} đồng\n"
            f"Tổng cộng: "
            f"{format_money(grand_total)} đồng",
        )

        open_file(output_path)

        return output_path

    except Exception as exc:  # noqa: BLE001
        messagebox.showerror(
            "Lỗi xuất Word", "Không thể xuất file Word.\n\n" + str(exc)
        )

        return None


# ============================================================
# SCROLLABLE FRAME
# ============================================================
# Không còn canvas lồng với main.py.
# Chỉ màn hình này quản lý scrollbar.
# ============================================================


class ScrollableFrame(ttk.Frame):
    def __init__(
        self,
        parent,
    ):
        super().__init__(parent)

        # ========================================================
        # CANVAS
        # ========================================================

        self.canvas = tk.Canvas(
            self,
            highlightthickness=0,
            borderwidth=0,
        )

        # ========================================================
        # SCROLLBAR DỌC
        # ========================================================

        self.v_scrollbar = ttk.Scrollbar(
            self,
            orient="vertical",
            command=self.canvas.yview,
        )

        # ========================================================
        # SCROLLBAR NGANG
        # ========================================================

        self.h_scrollbar = ttk.Scrollbar(
            self,
            orient="horizontal",
            command=self.canvas.xview,
        )

        # ========================================================
        # FRAME CHỨA TOÀN BỘ NỘI DUNG
        # ========================================================

        self.inner = ttk.Frame(self.canvas)

        self.window_id = self.canvas.create_window(
            (0, 0),
            window=self.inner,
            anchor="nw",
        )

        # ========================================================
        # KẾT NỐI SCROLLBAR
        # ========================================================

        self.canvas.configure(
            yscrollcommand=self.v_scrollbar.set,
            xscrollcommand=self.h_scrollbar.set,
        )

        # ========================================================
        # LAYOUT
        # ========================================================

        self.canvas.grid(
            row=0,
            column=0,
            sticky="nsew",
        )

        self.v_scrollbar.grid(
            row=0,
            column=1,
            sticky="ns",
        )

        self.h_scrollbar.grid(
            row=1,
            column=0,
            sticky="ew",
        )

        self.rowconfigure(
            0,
            weight=1,
        )

        self.columnconfigure(
            0,
            weight=1,
        )

        # ========================================================
        # CẬP NHẬT VÙNG SCROLL
        # ========================================================

        self.inner.bind(
            "<Configure>",
            self._update_scrollregion,
        )

        self.canvas.bind(
            "<Configure>",
            self._on_canvas_configure,
        )

        # ========================================================
        # MOUSE WHEEL
        # ========================================================

        self.canvas.bind(
            "<Enter>",
            self._bind_mousewheel,
        )

        self.canvas.bind(
            "<Leave>",
            self._unbind_mousewheel,
        )

        self.inner.bind(
            "<Enter>",
            self._bind_mousewheel,
        )

        self.inner.bind(
            "<Leave>",
            self._unbind_mousewheel,
        )

        # ========================================================
        # SHIFT + MOUSE WHEEL
        # -> CUỘN NGANG
        # ========================================================

        self.canvas.bind(
            "<Shift-MouseWheel>",
            self._on_shift_mousewheel,
        )

        self.inner.bind(
            "<Shift-MouseWheel>",
            self._on_shift_mousewheel,
        )

        # ========================================================
        # PHÍM MŨI TÊN
        # ========================================================

        self.canvas.bind(
            "<Left>",
            self._scroll_left,
        )

        self.canvas.bind(
            "<Right>",
            self._scroll_right,
        )

        # ========================================================
        # LUÔN BẮT ĐẦU Ở ĐẦU
        # ========================================================

        self.after_idle(self.scroll_to_top)

    # ========================================================
    # SCROLL REGION
    # ========================================================

    def _update_scrollregion(
        self,
        event=None,
    ):
        bbox = self.canvas.bbox("all")

        if bbox:
            self.canvas.configure(scrollregion=bbox)

    # ========================================================
    # CANVAS RESIZE
    # ========================================================

    def _on_canvas_configure(
        self,
        event,
    ):
        """
        Không ép chiều rộng inner bằng canvas.

        Đây là điểm quan trọng để scrollbar ngang hoạt động.
        """

        self._update_scrollregion()

    # ========================================================
    # BIND MOUSE WHEEL
    # ========================================================

    def _bind_mousewheel(
        self,
        event=None,
    ):
        self.canvas.bind_all(
            "<MouseWheel>",
            self._on_mousewheel,
        )

    # ========================================================
    # UNBIND MOUSE WHEEL
    # ========================================================

    def _unbind_mousewheel(
        self,
        event=None,
    ):
        self.canvas.unbind_all("<MouseWheel>")

    # ========================================================
    # CUỘN DỌC
    # ========================================================

    def _on_mousewheel(
        self,
        event,
    ):
        if event.delta:
            self.canvas.yview_scroll(
                int(-1 * (event.delta / 120)),
                "units",
            )

    # ========================================================
    # SHIFT + MOUSE WHEEL
    # -> CUỘN NGANG
    # ========================================================

    def _on_shift_mousewheel(
        self,
        event,
    ):
        if event.delta:
            self.canvas.xview_scroll(
                int(-1 * (event.delta / 120)),
                "units",
            )

        return "break"

    # ========================================================
    # CUỘN SANG TRÁI
    # ========================================================

    def _scroll_left(
        self,
        event=None,
    ):
        self.canvas.xview_scroll(
            -3,
            "units",
        )

        return "break"

    # ========================================================
    # CUỘN SANG PHẢI
    # ========================================================

    def _scroll_right(
        self,
        event=None,
    ):
        self.canvas.xview_scroll(
            3,
            "units",
        )

        return "break"

    # ========================================================
    # VỀ ĐẦU
    # ========================================================

    def scroll_to_top(
        self,
    ):
        self.canvas.yview_moveto(0)

        self.canvas.xview_moveto(0)


# ============================================================
# GIAO DIỆN CHÍNH
# ============================================================


class HopDongDuyetGiaFrame(ttk.Frame):
    def __init__(
        self,
        parent,
    ):
        super().__init__(parent)

        self.vendor_vars = []

        self.product_rows = []

        # self.create_style()

        self.create_ui()

        self.add_default_vendors()

        self.add_product()

        # QUAN TRỌNG:
        # Sau khi tạo toàn bộ nội dung,
        # đưa scrollbar về đầu.
        self.after_idle(self.scroll_to_top)

    # ========================================================
    # SCROLL
    # ========================================================

    def scroll_to_top(self):
        if hasattr(self, "main") and hasattr(
            self.main,
            "canvas",
        ):
            self.main.canvas.yview_moveto(0)

    # ========================================================
    # ========================================================
    # UI
    # ========================================================
    # ========================================================
    # NHẬP EXCEL
    # ========================================================

    def _excel_to_string(self, value):
        """Chuyển giá trị Excel thành chuỗi an toàn."""

        if value is None:
            return ""

        if isinstance(value, float) and value.is_integer():
            return str(int(value))

        return str(value).strip()

    def _read_excel_key_value_sheet(self, ws):
        """
        Đọc sheet dạng:

        bien | gia_tri | ghi_chu
        """

        data = {}

        for row in ws.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if not row:
                continue

            key = self._excel_to_string(row[0] if len(row) > 0 else "")

            value = self._excel_to_string(row[1] if len(row) > 1 else "")

            if key:
                data[key] = value

        return data

    def _read_excel_table(self, ws):
        """
        Đọc sheet dạng bảng:

        STT | name | quantity | ...
        """

        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            return []

        headers = [self._excel_to_string(value) for value in rows[0]]

        result = []

        for values in rows[1:]:
            if not values:
                continue

            item = {}
            has_value = False

            for index, header in enumerate(headers):
                if not header:
                    continue

                value = values[index] if index < len(values) else None

                value = self._excel_to_string(value)

                if value:
                    has_value = True

                item[header] = value

            if has_value:
                result.append(item)

        return result

    def _set_entry_value(self, widget, value):
        """Gán giá trị cho ttk.Entry."""

        widget.delete(
            0,
            tk.END,
        )

        widget.insert(
            0,
            str(value),
        )

    def _clear_imported_vendors(self):
        """
        Xóa các dòng nhà cung cấp hiện tại
        nhưng giữ lại dòng tiêu đề.
        """

        for widget in self.vendor_container.grid_slaves():
            info = widget.grid_info()

            try:
                row = int(info["row"])
            except Exception:
                continue

            if row >= 1:
                widget.destroy()

        self.vendor_vars = []

    def _clear_imported_products(self):
        """
        Xóa toàn bộ hạng mục hiện tại,
        giữ header và nút thêm.
        """

        for widget in self.product_container.grid_slaves():
            info = widget.grid_info()

            try:
                row = int(info["row"])
            except Exception:
                continue

            if row != 0 and row != 999:
                widget.destroy()

        self.product_rows = []

    def import_data_from_excel(self):
        """
        Nhập Excel bằng nút Chọn file Excel.
        """

        file_path = filedialog.askopenfilename(
            title="Chọn file Excel dữ liệu",
            filetypes=[
                (
                    "Excel Workbook",
                    "*.xlsx",
                ),
                (
                    "Excel Macro Workbook",
                    "*.xlsm",
                ),
            ],
        )

        if not file_path:
            return

        self.import_excel_file(file_path)

    def import_excel_file(self, file_path):
        """
        Đọc Excel và tự động điền TOÀN BỘ các field có trong giao diện.

        Excel gồm:
            - Thông tin chung
            - Nhà cung cấp
            - Hạng mục
            - Hợp đồng

        Nguyên tắc:
            1. Key trong Excel trùng tên biến/widget trên self
            -> tự động điền.
            2. Key trong Excel trùng key trong self.contract_vars
            -> tự động điền.
            3. Có alias cho các tên khác nhau giữa Excel và giao diện.
            4. Nhà cung cấp và Hạng mục được tạo lại từ Excel.
        """

        wb = None

        try:
            # ============================================================
            # 1. KIỂM TRA FILE
            # ============================================================

            if not file_path:
                return

            file_path = os.path.abspath(file_path)

            if not os.path.exists(file_path):
                raise FileNotFoundError(f"Không tìm thấy file:\n{file_path}")

            if not file_path.lower().endswith((".xlsx", ".xlsm")):
                raise ValueError("Chỉ hỗ trợ file Excel .xlsx hoặc .xlsm.")

            # ============================================================
            # 2. MỞ EXCEL
            # ============================================================

            wb = load_workbook(
                file_path,
                data_only=True,
            )

            # ============================================================
            # 3. KIỂM TRA SHEET
            # ============================================================

            required_sheets = [
                "Thông tin chung",
                "Nhà cung cấp",
                "Hạng mục",
                "Hợp đồng",
            ]

            missing_sheets = [
                sheet for sheet in required_sheets if sheet not in wb.sheetnames
            ]

            if missing_sheets:
                raise ValueError(
                    "File Excel thiếu sheet:\n\n"
                    + "\n".join(f"• {sheet}" for sheet in missing_sheets)
                    + "\n\n"
                    "Vui lòng sử dụng đúng file Excel mẫu."
                )

            # ============================================================
            # 4. HÀM PHỤ: ĐIỀN MỌI LOẠI FIELD
            # ============================================================

            def set_widget_value(widget, value):
                """
                Tự nhận diện:
                    - tk.StringVar
                    - tk.IntVar
                    - tk.DoubleVar
                    - tk.Text
                    - tk.Entry
                    - ttk.Entry
                    - ttk.Combobox
                """

                if widget is None:
                    return False

                if value is None:
                    value = ""

                value = str(value)

                # --------------------------------------------------------
                # StringVar / IntVar / DoubleVar...
                # --------------------------------------------------------

                if isinstance(widget, tk.Variable):
                    try:
                        widget.set(value)
                        return True
                    except Exception:
                        pass

                # --------------------------------------------------------
                # Text
                # --------------------------------------------------------

                if isinstance(widget, tk.Text):
                    try:
                        widget.delete("1.0", tk.END)
                        widget.insert("1.0", value)
                        return True
                    except Exception:
                        pass

                # --------------------------------------------------------
                # Entry / ttk.Entry
                # --------------------------------------------------------

                try:
                    widget.delete(0, tk.END)
                    widget.insert(0, value)
                    return True
                except Exception:
                    pass

                # --------------------------------------------------------
                # Combobox
                # --------------------------------------------------------

                try:
                    widget.set(value)
                    return True
                except Exception:
                    pass

                return False

            # ============================================================
            # 5. HÀM PHỤ: TÌM VÀ ĐIỀN FIELD
            # ============================================================

            def set_field(key, value):
                """
                Tìm field theo key trong toàn bộ giao diện.
                """

                if not key:
                    return False

                key = str(key).strip()

                # --------------------------------------------------------
                # BỎ QUA FIELD KHÔNG CÓ GIÁ TRỊ
                # --------------------------------------------------------

                if value is None:
                    return False

                value = str(value).strip()

                if value == "":
                    return False

                # ========================================================
                # 5.1. Tìm trực tiếp trên self
                #
                # Ví dụ:
                #
                # Excel:
                #     so_duyet
                #
                # Code:
                #     self.so_duyet
                # ========================================================

                if hasattr(self, key):
                    obj = getattr(self, key)

                    if set_widget_value(obj, value):
                        return True

                # ========================================================
                # 5.2. Tìm trong contract_vars
                # ========================================================

                contract_vars = getattr(self, "contract_vars", {})

                if key in contract_vars:
                    if set_widget_value(contract_vars[key], value):
                        return True

                # ========================================================
                # 5.3. ALIAS
                #
                # Excel có thể dùng tên khác với code.
                # ========================================================

                aliases = {
                    # ----------------------------------------------------
                    # Thông tin chung
                    # ----------------------------------------------------
                    "muc_dich_mua_sam": "muc_dich",
                    "muc_dich": "muc_dich",
                    "ten_duyet_gia": "ten_duyet_gia",
                    "so_duyet_gia": "so_duyet",
                    "so_duyet": "so_duyet",
                    "ngay_duyet": "ngay",
                    "thang_duyet": "thang",
                    "nam_duyet": "nam",
                    "dia_diem_duyet": "dia_diem",
                    "dia_diem": "dia_diem",
                    # ----------------------------------------------------
                    # Hợp đồng
                    # ----------------------------------------------------
                    "ma_hop_dong": "ma_hd",
                    "so_hop_dong": "ma_hd",
                    "ma_hd": "ma_hd",
                    "loai_hop_dong": "loai_hd",
                    "loai_hd": "loai_hd",
                }

                real_key = aliases.get(key)

                if real_key:
                    # Tìm trên self
                    if hasattr(self, real_key):
                        obj = getattr(self, real_key)

                        if set_widget_value(obj, value):
                            return True

                    # Tìm trong contract_vars
                    if real_key in contract_vars:
                        if set_widget_value(contract_vars[real_key], value):
                            return True

                return False

            # ============================================================
            # 6. ĐỌC THÔNG TIN CHUNG
            # ============================================================

            general = self._read_excel_key_value_sheet(wb["Thông tin chung"])

            imported_general = []
            ignored_general = []

            # ------------------------------------------------------------
            # QUAN TRỌNG:
            #
            # Không còn general_mapping cố định nữa.
            #
            # Excel có bao nhiêu key -> thử lấy bấy nhiêu.
            # ------------------------------------------------------------

            for key, value in general.items():
                if value is None:
                    continue

                if str(value).strip() == "":
                    continue

                # --------------------------------------------------------
                # Loại hợp đồng xử lý riêng ở bên dưới
                # --------------------------------------------------------

                if key == "loai_hd":
                    continue

                if set_field(key, value):
                    imported_general.append(key)

                else:
                    ignored_general.append(key)

            # ============================================================
            # 7. LOẠI HỢP ĐỒNG
            # ============================================================

            loai_hd = general.get("loai_hd", "")

            if loai_hd:
                loai_hd = str(loai_hd).strip()

                loai_hd_lower = loai_hd.lower()

                # --------------------------------------------------------
                # Chuẩn hóa
                # --------------------------------------------------------

                if loai_hd_lower in (
                    "dịch vụ",
                    "dich vu",
                    "dv",
                    "dịch-vụ",
                    "dich-vu",
                ):
                    loai_hd = "Dịch vụ"

                elif loai_hd_lower in (
                    "mua bán",
                    "mua ban",
                    "mb",
                    "mua-bán",
                    "mua-ban",
                ):
                    loai_hd = "Mua bán"

                else:
                    raise ValueError(
                        "Loại hợp đồng không hợp lệ:\n\n"
                        f"{loai_hd}\n\n"
                        "Chỉ chấp nhận:\n"
                        "• Dịch vụ\n"
                        "• Mua bán"
                    )

                try:
                    self.loai_hd.set(loai_hd)
                except Exception:
                    pass

                try:
                    self.change_contract_type()
                except Exception:
                    pass

                imported_general.append("loai_hd")

            # ============================================================
            # 8. NHÀ CUNG CẤP
            # ============================================================

            vendor_rows = self._read_excel_table(wb["Nhà cung cấp"])

            if len(vendor_rows) > MAX_VENDORS:
                raise ValueError(
                    f"Excel có {len(vendor_rows)} "
                    f"nhà cung cấp.\n\n"
                    f"Chương trình chỉ hỗ trợ tối đa "
                    f"{MAX_VENDORS} nhà cung cấp."
                )

            # ------------------------------------------------------------
            # Xóa NCC cũ
            # ------------------------------------------------------------

            self._clear_imported_vendors()

            # ------------------------------------------------------------
            # Tạo NCC mới
            # ------------------------------------------------------------

            for vendor in vendor_rows:
                vendor_data = {
                    "name": vendor.get("name", ""),
                    "nang_luc": vendor.get("nang_luc", ""),
                    "ky_thuat": vendor.get("ky_thuat", ""),
                    "dia_diem": vendor.get("dia_diem", ""),
                    "thoi_gian": vendor.get("thoi_gian", ""),
                    "cach_thanh_toan": vendor.get("cach_thanh_toan", ""),
                }

                # --------------------------------------------------------
                # Tạo dòng NCC
                # --------------------------------------------------------

                self.add_vendor(vendor_data)

                # --------------------------------------------------------
                # Nếu add_vendor tạo vendor_vars
                # thì tiếp tục kiểm tra các field còn lại.
                #
                # Điều này giúp sau này thêm field mới
                # vào vendor_vars mà không phải sửa import.
                # --------------------------------------------------------

                if self.vendor_vars:
                    current_vendor = self.vendor_vars[-1]

                    for key, value in vendor.items():
                        if not value:
                            continue

                        if key in current_vendor:
                            set_widget_value(current_vendor[key], value)

            # ============================================================
            # 9. HẠNG MỤC
            # ============================================================

            product_rows = self._read_excel_table(wb["Hạng mục"])

            # ------------------------------------------------------------
            # Xóa sản phẩm cũ
            # ------------------------------------------------------------

            self._clear_imported_products()

            # ------------------------------------------------------------
            # Tạo sản phẩm mới
            # ------------------------------------------------------------

            for product in product_rows:
                vars_ = {
                    "name": tk.StringVar(value=product.get("name", "")),
                    "quantity": tk.StringVar(value=product.get("quantity", "1")),
                    "unit": tk.StringVar(value=product.get("unit", "Cái")),
                    "price1": tk.StringVar(value=product.get("price1", "")),
                    "price2": tk.StringVar(value=product.get("price2", "")),
                    "price3": tk.StringVar(value=product.get("price3", "")),
                    "note": tk.StringVar(value=product.get("note", "")),
                }

                self._render_product_row(vars_)

            # ------------------------------------------------------------
            # Nếu Excel không có sản phẩm
            # ------------------------------------------------------------

            if not self.product_rows:
                self.add_product()

            # ============================================================
            # 10. HỢP ĐỒNG
            # ============================================================

            contract = self._read_excel_key_value_sheet(wb["Hợp đồng"])

            imported_contract = []
            ignored_contract = []

            # ------------------------------------------------------------
            # KHÔNG còn vòng:
            #
            # for key, var in self.contract_vars.items()
            #
            # đơn thuần nữa.
            #
            # Mà lấy TOÀN BỘ key Excel rồi tìm field tương ứng.
            # ------------------------------------------------------------

            for key, value in contract.items():
                if value is None:
                    continue

                if str(value).strip() == "":
                    continue

                if set_field(key, value):
                    imported_contract.append(key)

                else:
                    ignored_contract.append(key)

            # ============================================================
            # 11. ĐỒNG BỘ MÃ HỢP ĐỒNG
            #
            # Nếu cả 2 sheet đều có ma_hd thì ưu tiên Hợp đồng.
            # ============================================================

            if "ma_hd" in contract:
                self._set_entry_value(self.ma_hd, contract["ma_hd"])

            elif "ma_hd" in general:
                self._set_entry_value(self.ma_hd, general["ma_hd"])

            # ============================================================
            # 12. CẬP NHẬT NÚT / TỔNG
            # ============================================================

            try:
                self.update_add_button()
            except Exception:
                pass

            try:
                self.update_total()
            except Exception:
                pass

            # ============================================================
            # 13. CẬP NHẬT LABEL
            # ============================================================

            try:
                self.template_label.config(
                    text=("📊 Dữ liệu: " + os.path.basename(file_path))
                )

            except Exception:
                pass

            # ============================================================
            # 14. CUỘN LÊN ĐẦU
            # ============================================================

            try:
                self.after_idle(self.scroll_to_top)

            except Exception:
                pass

            # ============================================================
            # 15. ĐÓNG FILE
            # ============================================================

            try:
                wb.close()
            except Exception:
                pass

            wb = None

            # ============================================================
            # 16. THỐNG KÊ
            # ============================================================

            total_general = len(imported_general)

            total_contract = len(imported_contract)

            total_vendor = len(vendor_rows)

            total_product = len(product_rows)

            # ============================================================
            # 17. THÔNG BÁO
            # ============================================================

            message = (
                "ĐÃ NHẬP EXCEL THÀNH CÔNG\n\n"
                f"Thông tin chung: "
                f"{total_general} field\n"
                f"Hợp đồng: "
                f"{total_contract} field\n"
                f"Nhà cung cấp: "
                f"{total_vendor}\n"
                f"Hạng mục: "
                f"{total_product}"
            )

            # ------------------------------------------------------------
            # Các biến Excel có nhưng giao diện chưa có
            # ------------------------------------------------------------

            all_ignored = ignored_general + ignored_contract

            if all_ignored:
                message += (
                    "\n\n"
                    "⚠ Các biến Excel chưa tìm thấy "
                    "field tương ứng trên giao diện:\n\n"
                    + "\n".join(f"• {key}" for key in all_ignored)
                )

            messagebox.showinfo("Nhập Excel thành công", message)

        except Exception as exc:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

            messagebox.showerror(
                "Lỗi nhập Excel", "Không thể nhập dữ liệu:\n\n" + str(exc)
            )

    def on_excel_drop(self, event):
        try:
            files = self.tk.splitlist(event.data)

            if not files:
                return

            # Nếu kéo nhiều file thì chỉ lấy file đầu tiên
            file_path = files[0]

            # Loại bỏ dấu ngoặc nếu Windows trả về đường dẫn có ngoặc
            file_path = file_path.strip().strip('"')

            if not file_path.lower().endswith((".xlsx", ".xlsm")):
                messagebox.showerror(
                    "File không hợp lệ", "Vui lòng kéo thả file Excel .xlsx hoặc .xlsm."
                )
                return

            if not os.path.isfile(file_path):
                messagebox.showerror(
                    "Không tìm thấy file", f"Không tìm thấy:\n\n{file_path}"
                )
                return

            self.import_excel_file(file_path)

        except Exception as exc:
            traceback_text = str(exc)
            messagebox.showerror(
                "Lỗi kéo thả Excel", f"Không thể đọc file:\n\n{traceback_text}"
            )

    def create_ui(self):

        header = ttk.Frame(
            self,
            padding=(5, 5, 5, 3),
        )

        header.pack(
            fill="x",
            padx=5,
            pady=(0, 3),
        )

        ttk.Label(
            header,
            text=("PHÊ DUYỆT ĐÁNH GIÁ VÀ ĐỀ XUẤT"),
            style="Title.TLabel",
        ).pack(side="left")

        ttk.Button(
            header,
            text="📥 Nhập Excel",
            command=self.import_data_from_excel,
        ).pack(
            side="right",
            padx=(4, 0),
        )

        ttk.Button(
            header,
            text="📂 Chọn mẫu Excel",
            command=self.choose_template,
        ).pack(
            side="right",
            padx=(4, 0),
        )

        self.drop_area = tk.Label(
            self,
            text=("📥 KÉO THẢ FILE EXCEL VÀO ĐÂY\n(.xlsx / .xlsm)"),
            font=(
                "Arial",
                11,
                "bold",
            ),
            relief="groove",
            bd=2,
            padx=20,
            pady=12,
            cursor="hand2",
        )

        self.drop_area.pack(
            fill="x",
            padx=10,
            pady=(3, 6),
        )

        self.drop_area.drop_target_register(DND_FILES)

        self.drop_area.dnd_bind(
            "<<Drop>>",
            self.on_excel_drop,
        )

        self.drop_area.dnd_bind(
            "<<DragEnter>>",
            lambda e: self.drop_area.config(
                relief="sunken",
                text=("📥 THẢ FILE EXCEL TẠI ĐÂY"),
            ),
        )

        self.drop_area.dnd_bind(
            "<<DragLeave>>",
            lambda e: self.drop_area.config(
                relief="groove",
                text=("📥 KÉO THẢ FILE EXCEL VÀO ĐÂY\n(.xlsx / .xlsm)"),
            ),
        )

        ttk.Button(
            header,
            text="📂 Chọn file Word mẫu",
            command=self.choose_contract_template,
        ).pack(
            side="right",
            padx=(4, 0),
        )

        self.template_label = ttk.Label(
            header,
            text=(f"Excel: {os.path.abspath(TEMPLATE_FILE)}"),
        )

        self.template_label.pack(
            side="right",
            padx=20,
        )

        # ----------------------------------------------------
        # CHỈ CÓ MỘT SCROLLABLE FRAME
        # ----------------------------------------------------

        self.main = ScrollableFrame(self)

        self.main.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=(0, 5),
        )

        self.create_general_section()

        self.create_vendor_section()

        self.create_contract_section()

        self.create_product_section()

        self.create_action_section()

    # ========================================================
    # THÔNG TIN CHUNG
    # ========================================================

    def create_general_section(self):

        frame = ttk.LabelFrame(
            self.main.inner,
            text=" 1. Thông tin chung ",
            padding=6,
        )

        frame.pack(
            fill="x",
            pady=(0, 6),
        )

        for col in range(6):
            frame.columnconfigure(
                col,
                weight=1,
            )

        ttk.Label(
            frame,
            text="Số duyệt:",
            style="Header.TLabel",
        ).grid(
            row=0,
            column=0,
            sticky="w",
            padx=5,
            pady=5,
        )

        self.so_duyet = ttk.Entry(frame)

        self.so_duyet.grid(
            row=0,
            column=1,
            sticky="ew",
            padx=5,
            pady=5,
        )

        ttk.Label(
            frame,
            text="Ngày:",
            style="Header.TLabel",
        ).grid(
            row=0,
            column=2,
            sticky="w",
            padx=5,
            pady=5,
        )

        self.ngay = ttk.Entry(
            frame,
            width=8,
        )

        self.ngay.grid(
            row=0,
            column=3,
            sticky="ew",
            padx=5,
            pady=5,
        )

        ttk.Label(
            frame,
            text="Tháng:",
            style="Header.TLabel",
        ).grid(
            row=0,
            column=4,
            sticky="w",
            padx=5,
            pady=5,
        )

        self.thang = ttk.Entry(
            frame,
            width=8,
        )

        self.thang.grid(
            row=0,
            column=5,
            sticky="ew",
            padx=5,
            pady=5,
        )

        ttk.Label(
            frame,
            text="Năm:",
            style="Header.TLabel",
        ).grid(
            row=1,
            column=0,
            sticky="w",
            padx=5,
            pady=5,
        )

        self.nam = ttk.Entry(frame)

        self.nam.grid(
            row=1,
            column=1,
            sticky="ew",
            padx=5,
            pady=5,
        )

        now = datetime.now(VN_TZ)

        self.ngay.insert(
            0,
            f"{now.day:02d}",
        )

        self.thang.insert(
            0,
            f"{now.month:02d}",
        )

        self.nam.insert(
            0,
            str(now.year),
        )

        ttk.Label(
            frame,
            text="Địa điểm:",
            style="Header.TLabel",
        ).grid(
            row=1,
            column=2,
            sticky="w",
            padx=5,
            pady=5,
        )

        self.dia_diem = ttk.Entry(frame)

        self.dia_diem.insert(
            0,
            "Lào Cai",
        )

        self.dia_diem.grid(
            row=1,
            column=3,
            columnspan=3,
            sticky="ew",
            padx=5,
            pady=5,
        )

        ttk.Label(
            frame,
            text="Tên nội dung duyệt giá:",
            style="Header.TLabel",
        ).grid(
            row=2,
            column=0,
            sticky="nw",
            padx=5,
            pady=5,
        )

        self.ten_duyet_gia = tk.Text(
            frame,
            height=3,
            wrap="word",
        )

        self.ten_duyet_gia.grid(
            row=2,
            column=1,
            columnspan=5,
            sticky="ew",
            padx=5,
            pady=5,
        )

        ttk.Label(
            frame,
            text="Mục đích / yêu cầu mua sắm:",
            style="Header.TLabel",
        ).grid(
            row=3,
            column=0,
            sticky="nw",
            padx=5,
            pady=5,
        )

        self.muc_dich = tk.Text(
            frame,
            height=3,
            wrap="word",
        )

        self.muc_dich.grid(
            row=3,
            column=1,
            columnspan=5,
            sticky="ew",
            padx=5,
            pady=5,
        )

    # ========================================================
    # LOẠI HỢP ĐỒNG
    # ========================================================

    def change_contract_type(self):

        if self.loai_hd.get() == "Dịch vụ":
            template = DEFAULT_SERVICE_TEMPLATE
        else:
            template = DEFAULT_SALE_TEMPLATE

        self.template_file = template

        self.template_contract_file = template

        self.contract_template_label.config(
            text=(f"Mẫu Word: {os.path.abspath(template)}")
        )

    # ========================================================
    # NHÀ CUNG CẤP
    # ========================================================

    def create_vendor_section(self):

        frame = ttk.LabelFrame(
            self.main.inner,
            text=" 2. Nhà cung cấp ",
            padding=6,
        )

        frame.pack(
            fill="x",
            pady=(0, 6),
        )

        headers = [
            "STT",
            "Tên nhà cung cấp",
            "Năng lực",
            "Đáp ứng kỹ thuật",
            "Địa điểm giao hàng",
            "Thời gian thực hiện",
            "Thời hạn thanh toán",
        ]

        for col, text in enumerate(headers):
            ttk.Label(
                frame,
                text=text,
                style="Header.TLabel",
            ).grid(
                row=0,
                column=col,
                sticky="nsew",
                padx=3,
                pady=3,
            )

        for col in range(len(headers)):
            frame.columnconfigure(
                col,
                weight=1,
            )

            # STT
        frame.columnconfigure(0, weight=0, minsize=40)

        # Tên nhà cung cấp
        frame.columnconfigure(1, weight=2, minsize=140)

        # Năng lực
        frame.columnconfigure(2, weight=1, minsize=90)

        # Đáp ứng kỹ thuật
        frame.columnconfigure(3, weight=1, minsize=110)

        # Địa điểm
        frame.columnconfigure(4, weight=1, minsize=110)

        # Thời gian
        frame.columnconfigure(5, weight=2, minsize=150)

        # Thanh toán
        frame.columnconfigure(6, weight=2, minsize=150)

        self.vendor_container = frame

    def add_default_vendors(self):

        defaults = [
            {
                "name": "Nhà cung cấp 1",
                "nang_luc": "Đạt",
                "ky_thuat": "Đạt",
                "dia_diem": "Lào Cai",
                "thoi_gian": "Trong vòng 04 tuần kể từ ngày Bên A tạm ứng.",
                "cach_thanh_toan": "Theo thỏa thuận trong hợp đồng.",
            },
            {
                "name": "Nhà cung cấp 2",
                "nang_luc": "Đạt",
                "ky_thuat": "Đạt",
                "dia_diem": "Lào Cai",
                "thoi_gian": "Trong vòng 06 tuần kể từ ngày Bên A tạm ứng.",
                "cach_thanh_toan": "Theo thỏa thuận trong hợp đồng.",
            },
            {
                "name": "Nhà cung cấp 3",
                "nang_luc": "Đạt",
                "ky_thuat": "Đạt",
                "dia_diem": "Lào Cai",
                "thoi_gian": "Trong vòng 06 tuần kể từ ngày Bên A tạm ứng.",
                "cach_thanh_toan": "Theo thỏa thuận trong hợp đồng.",
            },
        ]

        for item in defaults:
            self.add_vendor(item)

    def add_vendor(
        self,
        data=None,
    ):

        if len(self.vendor_vars) >= MAX_VENDORS:
            messagebox.showinfo(
                "Thông báo",
                f"Mẫu Excel hiện tại hỗ trợ tối đa {MAX_VENDORS} nhà cung cấp.",
            )

            return

        data = data or {}

        row = len(self.vendor_vars) + 1

        variables = {
            "name": tk.StringVar(
                value=data.get(
                    "name",
                    "",
                )
            ),
            "nang_luc": tk.StringVar(
                value=data.get(
                    "nang_luc",
                    "Đạt",
                )
            ),
            "ky_thuat": tk.StringVar(
                value=data.get(
                    "ky_thuat",
                    "Đạt",
                )
            ),
            "dia_diem": tk.StringVar(
                value=data.get(
                    "dia_diem",
                    "",
                )
            ),
            "thoi_gian": tk.StringVar(
                value=data.get(
                    "thoi_gian",
                    "",
                )
            ),
            "cach_thanh_toan": tk.StringVar(
                value=data.get(
                    "cach_thanh_toan",
                    "",
                )
            ),
        }

        ttk.Label(
            self.vendor_container,
            text=str(row),
        ).grid(
            row=row,
            column=0,
            padx=3,
            pady=3,
        )

        for col, key in [
            (1, "name"),
            (2, "nang_luc"),
            (3, "ky_thuat"),
            (4, "dia_diem"),
            (5, "thoi_gian"),
            (6, "cach_thanh_toan"),
        ]:
            ttk.Entry(
                self.vendor_container,
                textvariable=variables[key],
            ).grid(
                row=row,
                column=col,
                sticky="ew",
                padx=3,
                pady=3,
            )

        self.vendor_vars.append(variables)

    # ========================================================
    # THÔNG TIN HỢP ĐỒNG
    # ========================================================

    def create_contract_section(self):

        frame = ttk.LabelFrame(
            self.main.inner,
            text=" 3. Thông tin hợp đồng / đối tác ",
            padding=6,
        )

        frame.pack(
            fill="x",
            pady=(0, 6),
        )

        for col in range(6):
            frame.columnconfigure(
                col,
                weight=1,
            )

        self.template_contract_file = DEFAULT_SERVICE_TEMPLATE

        self.template_file = DEFAULT_SERVICE_TEMPLATE

        self.contract_template_label = ttk.Label(
            frame,
            text=(f"Mẫu Word: {os.path.abspath(self.template_contract_file)}"),
        )

        self.contract_template_label.grid(
            row=0,
            column=0,
            columnspan=6,
            sticky="w",
            padx=5,
            pady=(0, 8),
        )

        ttk.Label(
            frame,
            text="Loại hợp đồng:",
            style="Header.TLabel",
        ).grid(
            row=1,
            column=0,
            sticky="w",
            padx=5,
            pady=5,
        )

        self.loai_hd = tk.StringVar(value="Dịch vụ")

        ttk.Radiobutton(
            frame,
            text="Dịch vụ",
            variable=self.loai_hd,
            value="Dịch vụ",
            command=self.change_contract_type,
        ).grid(
            row=1,
            column=1,
            sticky="w",
            padx=5,
        )

        ttk.Radiobutton(
            frame,
            text="Mua bán",
            variable=self.loai_hd,
            value="Mua bán",
            command=self.change_contract_type,
        ).grid(
            row=1,
            column=2,
            sticky="w",
            padx=5,
        )

        ttk.Label(
            frame,
            text="Mã hợp đồng:",
            style="Header.TLabel",
        ).grid(
            row=1,
            column=3,
            sticky="w",
            padx=5,
        )

        self.ma_hd = ttk.Entry(frame)

        self.ma_hd.grid(
            row=1,
            column=4,
            columnspan=2,
            sticky="ew",
            padx=5,
        )

        fields = [
            (
                "Mã đối tác",
                "ma_doi_tac",
            ),
            (
                "Đối tác lập hợp đồng",
                "ten_doi_tac",
            ),
            (
                "Đại diện",
                "dai_dien_doi_tac",
            ),
            (
                "Chức vụ",
                "chuc_vu_dai_dien_doi_tac",
            ),
            (
                "Địa chỉ đối tác",
                "dia_chi_doi_tac",
            ),
            (
                "Điện thoại",
                "sdt_doi_tac",
            ),
            (
                "Số tài khoản",
                "stk_doi_tac",
            ),
            (
                "Ngân hàng",
                "ngan_hang_doi_tac",
            ),
            (
                "Chi nhánh ngân hàng",
                "chi_nhanh_ngan_hang_doi_tac",
            ),
            (
                "Mã số thuế",
                "ma_so_thue_doi_tac",
            ),
            (
                "Thời gian giao hàng/thực hiện",
                "thoi_gian_giao_hang",
            ),
            (
                "Địa điểm giao hàng/thực hiện",
                "dia_chi_giao_hang",
            ),
            (
                "Địa điểm bảo hành",
                "dia_diem_bao_hanh",
            ),
            (
                "Thời gian bảo hành",
                "thoi_gian_bao_hanh",
            ),
            (
                "Tòa án",
                "toa_an",
            ),
        ]

        self.contract_vars = {}

        for i, (
            label,
            key,
        ) in enumerate(
            fields,
            start=2,
        ):
            r = i

            c = 0 if i % 2 == 0 else 3

            value_col = c + 1

            ttk.Label(
                frame,
                text=label + ":",
                style="Header.TLabel",
            ).grid(
                row=r,
                column=c,
                sticky="w",
                padx=5,
                pady=3,
            )

            var = tk.StringVar()

            self.contract_vars[key] = var

            ttk.Entry(
                frame,
                textvariable=var,
            ).grid(
                row=r,
                column=value_col,
                columnspan=2,
                sticky="ew",
                padx=5,
                pady=3,
            )

    # ========================================================
    # SẢN PHẨM
    # ========================================================

    def create_product_section(self):

        frame = ttk.LabelFrame(
            self.main.inner,
            text=" 4. Hạng mục / vật tư ",
            padding=6,
        )

        frame.pack(
            fill="x",
            pady=(0, 6),
        )

        headers = [
            "STT",
            "Tên hạng mục",
            "Số lượng",
            "ĐVT",
            "Giá NCC 1",
            "Giá NCC 2",
            "Giá NCC 3",
            "Ghi chú",
            "",
        ]

        for col, text in enumerate(headers):
            ttk.Label(
                frame,
                text=text,
                style="Header.TLabel",
            ).grid(
                row=0,
                column=col,
                sticky="nsew",
                padx=3,
                pady=3,
            )

        widths = [
            5,
            28,
            10,
            8,
            15,
            15,
            15,
            20,
            8,
        ]

        for col, weight in enumerate(widths):
            frame.columnconfigure(
                col,
                weight=weight,
            )

        self.product_container = frame

        ttk.Button(
            frame,
            text="➕ Thêm hạng mục",
            command=self.add_product,
        ).grid(
            row=999,
            column=0,
            columnspan=9,
            sticky="w",
            padx=3,
            pady=8,
        )

    def add_product(self):

        if len(self.product_rows) >= 100:
            messagebox.showwarning("Giới hạn", "Không nên có quá 100 hạng mục.")

            return

        vars_ = {
            "name": tk.StringVar(),
            "quantity": tk.StringVar(value="1"),
            "unit": tk.StringVar(value="Cái"),
            "price1": tk.StringVar(),
            "price2": tk.StringVar(),
            "price3": tk.StringVar(),
            "note": tk.StringVar(),
        }

        self._render_product_row(vars_)

        self.update_add_button()

    def _render_product_row(
        self,
        vars_,
    ):

        row = len(self.product_rows) + 1

        ttk.Label(
            self.product_container,
            text=str(row),
        ).grid(
            row=row,
            column=0,
            padx=3,
            pady=3,
        )

        ttk.Entry(
            self.product_container,
            textvariable=vars_["name"],
        ).grid(
            row=row,
            column=1,
            sticky="ew",
            padx=3,
            pady=3,
        )

        ttk.Entry(
            self.product_container,
            textvariable=vars_["quantity"],
        ).grid(
            row=row,
            column=2,
            sticky="ew",
            padx=3,
            pady=3,
        )

        ttk.Entry(
            self.product_container,
            textvariable=vars_["unit"],
        ).grid(
            row=row,
            column=3,
            sticky="ew",
            padx=3,
            pady=3,
        )

        for col, key in [
            (4, "price1"),
            (5, "price2"),
            (6, "price3"),
        ]:
            ttk.Entry(
                self.product_container,
                textvariable=vars_[key],
            ).grid(
                row=row,
                column=col,
                sticky="ew",
                padx=3,
                pady=3,
            )

        ttk.Entry(
            self.product_container,
            textvariable=vars_["note"],
        ).grid(
            row=row,
            column=7,
            sticky="ew",
            padx=3,
            pady=3,
        )

        if row > 1:
            ttk.Button(
                self.product_container,
                text="Xóa",
                command=lambda r=row: self.remove_product(r),
            ).grid(
                row=row,
                column=8,
                padx=3,
                pady=3,
            )

        self.product_rows.append(vars_)

    def update_add_button(self):
        """
        Cập nhật trạng thái nút Thêm hạng mục.
        """

        for widget in self.product_container.grid_slaves(
            row=999,
            column=0,
        ):
            widget.destroy()

        if len(self.product_rows) < 100:
            ttk.Button(
                self.product_container,
                text="➕ Thêm hạng mục",
                command=self.add_product,
            ).grid(
                row=999,
                column=0,
                columnspan=9,
                sticky="w",
                padx=3,
                pady=8,
            )

    # ========================================================
    # ACTION
    # ========================================================

    def create_action_section(self):

        frame = ttk.Frame(
            self.main.inner,
            padding=10,
        )

        frame.pack(
            fill="x",
            pady=10,
        )

        ttk.Button(
            frame,
            text="💾 XUẤT EXCEL",
            style="Action.TButton",
            command=self.export,
        ).pack(
            side="left",
            padx=(4, 0),
        )

        ttk.Button(
            frame,
            text="🔍 Xem dữ liệu",
            style="Action.TButton",
            command=self.preview_data,
        ).pack(
            side="left",
            padx=(4, 0),
        )

        ttk.Button(
            frame,
            text="📁 Mở thư mục output",
            style="Action.TButton",
            command=self.open_output_dir,
        ).pack(
            side="left",
            padx=(4, 0),
        )

        ttk.Button(
            frame,
            text="📄 XUẤT HỢP ĐỒNG WORD",
            style="Action.TButton",
            command=self.export_contract,
        ).pack(
            side="left",
            padx=(4, 0),
        )

    # ========================================================
    # DATA
    # ========================================================

    def get_data(self):

        vendors = []

        for vendor in self.vendor_vars:
            name = vendor["name"].get().strip()

            if not name:
                continue

            vendors.append(
                {
                    "name": name,
                    "nang_luc": vendor["nang_luc"].get(),
                    "ky_thuat": vendor["ky_thuat"].get(),
                    "dia_diem": vendor["dia_diem"].get(),
                    "thoi_gian": vendor["thoi_gian"].get(),
                    "cach_thanh_toan": vendor["cach_thanh_toan"].get(),
                }
            )

        products = []

        for index, row in enumerate(
            self.product_rows,
            start=1,
        ):
            name = row["name"].get().strip()

            if not name:
                continue

            products.append(
                {
                    "id": index,
                    "name": name,
                    "quantity": row["quantity"].get().strip(),
                    "unit": row["unit"].get().strip(),
                    "prices": [
                        row["price1"].get().strip(),
                        row["price2"].get().strip(),
                        row["price3"].get().strip(),
                    ],
                    "note": row["note"].get().strip(),
                }
            )

        for product_index, product in enumerate(
            products,
            start=1,
        ):
            for i in range(len(vendors)):
                price = product["prices"][i].strip()

                if not price:
                    continue

                try:
                    parse_number(price)
                except Exception:  # noqa: BLE001
                    raise ValueError(
                        f"Hạng mục {product_index}, "
                        f"giá NCC {i + 1} không hợp lệ:\n"
                        f"{price}"
                    )
        data = {
            "so_duyet": self.so_duyet.get().strip(),
            "ngay": self.ngay.get().strip(),
            "thang": self.thang.get().strip(),
            "nam": self.nam.get().strip(),
            "ten_duyet_gia": self.ten_duyet_gia.get(
                "1.0",
                "end",
            ).strip(),
            "muc_dich_mua_sam": self.muc_dich.get(
                "1.0",
                "end",
            ).strip(),
            "dia_diem": self.dia_diem.get().strip(),
            "ma_hd": self.ma_hd.get().strip(),
            "loai_hd": self.loai_hd.get(),
            "vendors": vendors,
            "products": products,
        }

        data.update({key: var.get().strip() for key, var in self.contract_vars.items()})

        return data

    # ========================================================
    # EXPORT EXCEL
    # ========================================================

    def export(self):

        try:
            data = self.get_data()

            if not data["vendors"]:
                raise ValueError("Chưa nhập nhà cung cấp.")

            if len(data["vendors"]) > MAX_VENDORS:
                raise ValueError(f"Tối đa {MAX_VENDORS} nhà cung cấp.")

            if not data["products"]:
                raise ValueError("Chưa nhập hạng mục.")

            for index, product in enumerate(
                data["products"],
                start=1,
            ):
                for i in range(len(data["vendors"])):
                    if not product["prices"][i].strip():
                        raise ValueError(
                            f"Hạng mục {index} chưa nhập giá của NCC {i + 1}."
                        )

            os.makedirs(
                DEFAULT_OUTPUT_DIR,
                exist_ok=True,
            )

            default_name = (
                "duyet_gia_" + datetime.now(VN_TZ).strftime("%Y%m%d_%H%M%S") + ".xlsx"
            )

            output_path = filedialog.asksaveasfilename(
                title="Lưu file Excel",
                initialdir=os.path.abspath(DEFAULT_OUTPUT_DIR),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[
                    (
                        "Excel Workbook",
                        "*.xlsx",
                    )
                ],
            )

            if not output_path:
                return

            totals = []

            for vendor_index in range(len(data["vendors"])):
                total = sum(
                    parse_number(p["prices"][vendor_index])
                    * parse_number(p["quantity"])
                    for p in data["products"]
                )

                totals.append(total)

            data["selected_vendor_index"] = totals.index(min(totals))

            result = export_duyet_gia(
                data=data,
                template_path=TEMPLATE_FILE,
                output_path=output_path,
            )

            messagebox.showinfo(
                "Xuất Excel thành công",
                "Đã tạo file:\n\n"
                f"{result['output_path']}\n\n"
                "Nhà cung cấp được chọn "
                "theo tổng giá:\n"
                f"{result['selected_vendor']}\n\n"
                "Tổng giá chưa VAT:\n"
                f"{format_money(result['selected_total'])} đồng",
            )

            open_file(output_path)

        except Exception as exc:  # noqa: BLE001
            messagebox.showerror(
                "Lỗi",
                str(exc),
            )

    # ========================================================
    # EXPORT WORD
    # ========================================================

    def export_contract(self):

        try:
            data = self.get_data()

            if not data["vendors"]:
                raise ValueError("Chưa nhập nhà cung cấp.")

            if not data["products"]:
                raise ValueError("Chưa nhập hạng mục.")

            required = [
                (
                    "Mã hợp đồng",
                    "ma_hd",
                ),
                (
                    "Mã đối tác",
                    "ma_doi_tac",
                ),
                (
                    "Tên đối tác",
                    "ten_doi_tac",
                ),
                (
                    "Đại diện đối tác",
                    "dai_dien_doi_tac",
                ),
                (
                    "Chức vụ đại diện",
                    "chuc_vu_dai_dien_doi_tac",
                ),
                (
                    "Địa chỉ đối tác",
                    "dia_chi_doi_tac",
                ),
                (
                    "Thời gian giao hàng/thực hiện",
                    "thoi_gian_giao_hang",
                ),
                (
                    "Địa điểm giao hàng/thực hiện",
                    "dia_chi_giao_hang",
                ),
            ]

            missing = [label for label, key in required if not data.get(key)]

            if missing:
                raise ValueError("Chưa nhập: " + ", ".join(missing))

            totals = [
                sum(
                    parse_number(p["prices"][i]) * parse_number(p["quantity"])
                    for p in data["products"]
                )
                for i in range(len(data["vendors"]))
            ]

            data["selected_vendor_index"] = totals.index(min(totals))

            template = (
                DEFAULT_SERVICE_TEMPLATE
                if data["loai_hd"] == "Dịch vụ"
                else DEFAULT_SALE_TEMPLATE
            )

            if self.template_contract_file:
                template = self.template_contract_file

            os.makedirs(
                DEFAULT_CONTRACT_OUTPUT_DIR,
                exist_ok=True,
            )

            suffix = "HĐDV" if data["loai_hd"] == "Dịch vụ" else "HĐMB"

            default_name = (
                f"{suffix}_"
                f"{data['ma_hd']}_"
                f"{data['ma_doi_tac']}_"
                f"{datetime.now(VN_TZ).strftime('%Y%m%d_%H%M%S')}"
                ".docx"
            )

            output_path = filedialog.asksaveasfilename(
                title="Lưu hợp đồng Word",
                initialdir=os.path.abspath(DEFAULT_CONTRACT_OUTPUT_DIR),
                initialfile=default_name,
                defaultextension=".docx",
                filetypes=[
                    (
                        "Word Document",
                        "*.docx",
                    )
                ],
            )

            if not output_path:
                return

            generate_contract_docx(self)

        except Exception as exc:  # noqa: BLE001
            messagebox.showerror(
                "Lỗi xuất hợp đồng",
                str(exc),
            )

    # ========================================================
    # PREVIEW
    # ========================================================

    def preview_data(self):

        try:
            data = self.get_data()

            text = []

            text.append("===== NHÀ CUNG CẤP =====")

            for i, vendor in enumerate(
                data["vendors"],
                start=1,
            ):
                text.append(f"{i}. {vendor['name']}")

            text.append("")

            text.append("===== HẠNG MỤC =====")

            totals = [0] * len(data["vendors"])

            for i, product in enumerate(
                data["products"],
                start=1,
            ):
                text.append(
                    f"{i}. {product['name']} - {product['quantity']} {product['unit']}"
                )

                for vendor_index in range(len(data["vendors"])):
                    price = parse_number(product["prices"][vendor_index])

                    quantity = parse_number(product["quantity"])

                    totals[vendor_index] += price * quantity

                    text.append(f"   NCC {vendor_index + 1}: {format_money(price)}")

            if totals:
                selected_index = totals.index(min(totals))

                text.append("")

                text.append("===== TỔNG =====")

                for i, total in enumerate(
                    totals,
                    start=1,
                ):
                    text.append(f"NCC {i}: {format_money(total)} đồng")

                text.append("")

                text.append(
                    f"Đề xuất theo tổng giá: {data['vendors'][selected_index]['name']}"
                )

            messagebox.showinfo(
                "Xem dữ liệu",
                "\n".join(text),
            )

        except Exception as exc:
            messagebox.showerror(
                "Lỗi dữ liệu",
                str(exc),
            )

    # ========================================================
    # TEMPLATE
    # ========================================================

    def choose_template(self):

        global TEMPLATE_FILE

        path = filedialog.askopenfilename(
            title="Chọn file Excel mẫu",
            filetypes=[
                (
                    "Excel Workbook",
                    "*.xlsx",
                )
            ],
        )

        if not path:
            return

        TEMPLATE_FILE = path

        self.template_label.config(text=(f"Excel: {os.path.abspath(TEMPLATE_FILE)}"))

    def choose_contract_template(self):

        file_path = filedialog.askopenfilename(
            title="Chọn file Word template",
            filetypes=[
                (
                    "Word files",
                    "*.docx",
                ),
                (
                    "All files",
                    "*.*",
                ),
            ],
        )

        if not file_path:
            return

        self.template_contract_file = file_path

        self.template_file = file_path

        self.contract_template_label.config(
            text=(f"Mẫu Word: {os.path.abspath(file_path)}")
        )

    # ========================================================
    # OUTPUT
    # ========================================================

    def open_output_dir(self):
        path = os.path.abspath(DEFAULT_OUTPUT_DIR)

        os.makedirs(
            path,
            exist_ok=True,
        )
        open_file(path)
