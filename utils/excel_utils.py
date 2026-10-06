import os
import zipfile
import copy
from copy import deepcopy
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


def copy_row_style(
    ws,
    source_row,
    target_row,
    min_col=1,
    max_col=10,
):
    """Copy toàn bộ style của một dòng."""

    for col in range(
        min_col,
        max_col + 1,
    ):
        source = ws.cell(
            source_row,
            col,
        )

        target = ws.cell(
            target_row,
            col,
        )

        if source.has_style:
            target._style = copy.copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.font:
            target.font = copy.copy(source.font)

        if source.fill:
            target.fill = copy.copy(source.fill)

        if source.border:
            target.border = copy.copy(source.border)

        if source.alignment:
            target.alignment = copy.copy(source.alignment)

        if source.protection:
            target.protection = copy.copy(source.protection)

    if ws.row_dimensions[source_row].height is not None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def clear_row(
    ws,
    row,
    min_col=1,
    max_col=10,
):
    for col in range(
        min_col,
        max_col + 1,
    ):
        cell = ws.cell(
            row,
            col,
        )

        if isinstance(
            cell,
            MergedCell,
        ):
            continue

        cell.value = None


def replace_text_in_sheet(
    ws,
    old,
    new,
):
    for row in ws.iter_rows():
        for cell in row:
            if (
                isinstance(
                    cell.value,
                    str,
                )
                and old in cell.value
            ):
                cell.value = cell.value.replace(
                    old,
                    new,
                )


def _excel_text_width(
    ws,
    row,
    col,
):
    return float(ws.column_dimensions[get_column_letter(col)].width or 8.43)


def _merged_width(
    ws,
    row,
    col,
):
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = rng.bounds

        if min_row <= row <= max_row and min_col <= col <= max_col:
            if row == min_row and col == min_col:
                return sum(
                    _excel_text_width(
                        ws,
                        row,
                        c,
                    )
                    for c in range(
                        min_col,
                        max_col + 1,
                    )
                )

            return 0

    return _excel_text_width(
        ws,
        row,
        col,
    )


def auto_fit_row_heights(
    ws,
    min_row=1,
    max_row=None,
    min_height=15.0,
    max_height=409.0,
):
    if max_row is None:
        max_row = ws.max_row

    for row in range(
        min_row,
        max_row + 1,
    ):
        required_height = min_height

        for col in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row,
                col,
            )

            if isinstance(
                cell,
                MergedCell,
            ):
                continue

            value = cell.value

            if (
                value is None
                or not isinstance(
                    value,
                    str,
                )
                or not value.strip()
            ):
                continue

            alignment = copy.copy(cell.alignment)

            alignment.wrap_text = True

            cell.alignment = alignment

            width = _merged_width(
                ws,
                row,
                col,
            )

            if width <= 0:
                continue

            font_size = float(cell.font.sz or 11)

            font_factor = max(
                0.72,
                min(
                    1.0,
                    11.0 / font_size,
                ),
            )

            chars_per_line = max(
                5,
                int(width * 0.95 * font_factor),
            )

            line_count = 0

            for raw_line in (
                str(value)
                .replace(
                    "\r\n",
                    "\n",
                )
                .replace(
                    "\r",
                    "\n",
                )
                .split("\n")
            ):
                raw_line = raw_line.expandtabs(4)

                if not raw_line:
                    line_count += 1
                    continue

                line_count += max(
                    1,
                    (len(raw_line) + chars_per_line - 1) // chars_per_line,
                )

            line_height = max(
                13.5,
                font_size * 1.30,
            )

            cell_height = line_count * line_height + 4

            required_height = max(
                required_height,
                cell_height,
            )

        if required_height > min_height:
            ws.row_dimensions[row].height = min(
                required_height,
                max_height,
            )

        elif ws.row_dimensions[row].height is None:
            ws.row_dimensions[row].height = min_height


def prepare_product_area(
    ws,
    product_count,
):
    product_template_row = 18
    first_evaluation_row = 19

    if product_count <= 0:
        product_count = 1

    delta = product_count - 1

    if delta <= 0:
        return product_template_row

    merge_ranges_to_shift = []

    for merge_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = merge_range.bounds

        if min_row >= first_evaluation_row:
            merge_ranges_to_shift.append(
                (
                    min_col,
                    min_row,
                    max_col,
                    max_row,
                )
            )

    for (
        min_col,
        min_row,
        max_col,
        max_row,
    ) in merge_ranges_to_shift:
        ws.unmerge_cells(
            f"{get_column_letter(min_col)}"
            f"{min_row}:"
            f"{get_column_letter(max_col)}"
            f"{max_row}"
        )

    ws.insert_rows(
        first_evaluation_row,
        delta,
    )

    for row in range(
        first_evaluation_row,
        first_evaluation_row + delta,
    ):
        copy_row_style(
            ws,
            product_template_row,
            row,
            min_col=1,
            max_col=10,
        )

    for (
        min_col,
        min_row,
        max_col,
        max_row,
    ) in merge_ranges_to_shift:
        shifted_min_row = min_row + delta

        shifted_max_row = max_row + delta

        ws.merge_cells(
            f"{get_column_letter(min_col)}"
            f"{shifted_min_row}:"
            f"{get_column_letter(max_col)}"
            f"{shifted_max_row}"
        )

    return product_template_row


def restore_template_drawing(
    template_path,
    output_path,
):
    import xml.etree.ElementTree as ET

    if not os.path.exists(template_path) or not os.path.exists(output_path):
        return

    drawing_name = "xl/drawings/drawing1.xml"

    rels_name = "xl/worksheets/_rels/sheet1.xml.rels"

    sheet_name = "xl/worksheets/sheet1.xml"

    content_types_name = "[Content_Types].xml"

    NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    NS_PKGREL = "http://schemas.openxmlformats.org/package/2006/relationships"

    NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"

    ET.register_namespace(
        "",
        NS_MAIN,
    )

    ET.register_namespace(
        "r",
        NS_REL,
    )

    with zipfile.ZipFile(
        template_path,
        "r",
    ) as template_zip:
        if drawing_name not in template_zip.namelist():
            return

        drawing_data = template_zip.read(drawing_name)

    temp_path = output_path + ".tmp_restore.xlsx"

    with (
        zipfile.ZipFile(
            output_path,
            "r",
        ) as src_zip,
        zipfile.ZipFile(
            temp_path,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as dst_zip,
    ):
        output_names = set(src_zip.namelist())

        if rels_name in output_names:
            rel_root = ET.fromstring(src_zip.read(rels_name))
        else:
            rel_root = ET.Element(f"{{{NS_PKGREL}}}Relationships")

        drawing_rel = None

        for rel in rel_root.findall(f"{{{NS_PKGREL}}}Relationship"):
            if rel.get(
                "Type",
                "",
            ).endswith("/drawing"):
                drawing_rel = rel
                break

        if drawing_rel is None:
            used_ids = {
                rel.get("Id")
                for rel in rel_root.findall(f"{{{NS_PKGREL}}}Relationship")
            }

            n = 1

            while f"rId{n}" in used_ids:
                n += 1

            drawing_rel = ET.SubElement(
                rel_root,
                f"{{{NS_PKGREL}}}Relationship",
                {
                    "Id": f"rId{n}",
                    "Type": (
                        "http://schemas.openxmlformats.org/"
                        "officeDocument/2006/"
                        "relationships/drawing"
                    ),
                    "Target": ("../drawings/drawing1.xml"),
                },
            )

        else:
            drawing_rel.set(
                "Target",
                "../drawings/drawing1.xml",
            )

        drawing_rid = drawing_rel.get("Id")

        rels_data = ET.tostring(
            rel_root,
            encoding="utf-8",
            xml_declaration=True,
        )

        sheet_root = ET.fromstring(src_zip.read(sheet_name))

        drawing_element = sheet_root.find(f"{{{NS_MAIN}}}drawing")

        if drawing_element is None:
            drawing_element = ET.Element(
                f"{{{NS_MAIN}}}drawing",
                {f"{{{NS_REL}}}id": drawing_rid},
            )

            children = list(sheet_root)

            insert_index = len(children)

            for i, child in enumerate(children):
                if child.tag in (
                    f"{{{NS_MAIN}}}pageMargins",
                    f"{{{NS_MAIN}}}pageSetup",
                    f"{{{NS_MAIN}}}headerFooter",
                    f"{{{NS_MAIN}}}rowBreaks",
                    f"{{{NS_MAIN}}}colBreaks",
                ):
                    insert_index = i + 1

            sheet_root.insert(
                insert_index,
                drawing_element,
            )

        else:
            drawing_element.set(
                f"{{{NS_REL}}}id",
                drawing_rid,
            )

        sheet_data = ET.tostring(
            sheet_root,
            encoding="utf-8",
            xml_declaration=True,
        )

        content_types_root = ET.fromstring(src_zip.read(content_types_name))

        has_drawing_content_type = any(
            item.get("PartName") == "/xl/drawings/drawing1.xml"
            for item in content_types_root.findall(f"{{{NS_CT}}}Override")
        )

        if not has_drawing_content_type:
            ET.SubElement(
                content_types_root,
                f"{{{NS_CT}}}Override",
                {
                    "PartName": "/xl/drawings/drawing1.xml",
                    "ContentType": "application/"
                    "vnd.openxmlformats-officedocument."
                    "drawing+xml",
                },
            )

        content_types_data = ET.tostring(
            content_types_root,
            encoding="utf-8",
            xml_declaration=True,
        )

        for item in src_zip.infolist():
            name = item.filename

            if name == rels_name:
                dst_zip.writestr(
                    item,
                    rels_data,
                )

            elif name == sheet_name:
                dst_zip.writestr(
                    item,
                    sheet_data,
                )

            elif name == content_types_name:
                dst_zip.writestr(
                    item,
                    content_types_data,
                )

            elif name == drawing_name:
                dst_zip.writestr(
                    item,
                    drawing_data,
                )

            else:
                dst_zip.writestr(
                    item,
                    src_zip.read(name),
                )

        if drawing_name not in output_names:
            drawing_info = zipfile.ZipInfo(drawing_name)

            drawing_info.compress_type = zipfile.ZIP_DEFLATED

            dst_zip.writestr(
                drawing_info,
                drawing_data,
            )

    os.replace(
        temp_path,
        output_path,
    )
