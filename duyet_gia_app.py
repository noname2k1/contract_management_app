import os
import re
import sys
import copy
import subprocess
import zipfile
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from utils import (
    parse_number,
    format_money,
    number_to_vietnamese_words,
    generate_docx,
    flatten_jobs,
    calculate_service_totals,
)


# ============================================================
# CẤU HÌNH
# ============================================================

TEMPLATE_FILE = "./templates/duyet_gia.xlsx"
DEFAULT_OUTPUT_DIR = "./outputs"
DEFAULT_SERVICE_TEMPLATE = "./templates/TB1.docx"
DEFAULT_SALE_TEMPLATE = "./templates/TB1.docx"

MAX_VENDORS = 3


# ============================================================
# HÀM TIỆN ÍCH
# ============================================================


def copy_row_style(ws, source_row, target_row, min_col=1, max_col=10):
    """Copy style, number format, alignment, border, fill... từ dòng mẫu."""
    for col in range(min_col, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)

        if source.has_style:
            target._style = copy.copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.font:
            target.font = copy.copy(source.font)

        if source.fill:
            target.fill = copy.copy(source.fill)

        if source.border:
            target.border = copy.copy(source.border)

        if source.alignment:
            target.alignment = copy.copy(source.alignment)

        if source.protection:
            target.protection = copy.copy(source.protection)

    if ws.row_dimensions[source_row].height is not None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def clear_row(ws, row, min_col=1, max_col=10):
    # Excel mẫu có nhiều vùng merge (A14:J14, A15:J15...).
    # Các ô phụ trong vùng merge là MergedCell và không cho phép
    # gán .value. Chỉ xóa ô thực sự (Cell).
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = None


def replace_text_in_sheet(ws, old, new):
    """Thay text trong toàn bộ sheet."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and old in cell.value:
                cell.value = cell.value.replace(old, new)


def open_file(path):
    path = os.path.abspath(path)

    if not os.path.exists(path):
        return

    try:
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
    except Exception as exc:
        messagebox.showwarning(
            "Không mở được file", f"File đã được tạo nhưng không thể tự mở:\n{exc}"
        )


# ============================================================
# XỬ LÝ EXCEL
# ============================================================


def _excel_text_width(ws, row, col):
    """Ước lượng độ rộng hiển thị của ô theo chiều rộng cột."""
    return float(ws.column_dimensions[get_column_letter(col)].width or 8.43)


def _merged_width(ws, row, col):
    """
    Nếu ô là ô đầu của vùng merge ngang thì trả về tổng độ rộng các cột.
    Nếu không phải ô đầu merge thì trả về 0 để không tính trùng.
    """
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = rng.bounds
        if min_row <= row <= max_row and min_col <= col <= max_col:
            if row == min_row and col == min_col:
                return sum(
                    _excel_text_width(ws, row, c) for c in range(min_col, max_col + 1)
                )
            return 0
    return _excel_text_width(ws, row, col)


def auto_fit_row_heights(
    ws, min_row=1, max_row=None, min_height=15.0, max_height=409.0
):
    """
    Tự động tăng chiều cao dòng theo nội dung, kể cả ô Merge.

    openpyxl không có AutoFit Row Height chuẩn như Excel, đặc biệt với
    ô Merge, nên hàm này ước lượng số dòng dựa trên nội dung, độ rộng
    cột/vùng Merge và cỡ chữ.
    """
    if max_row is None:
        max_row = ws.max_row

    for row in range(min_row, max_row + 1):
        required_height = min_height

        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)

            # Ô phụ của vùng Merge không được gán giá trị/format riêng.
            if isinstance(cell, MergedCell):
                continue

            value = cell.value
            if value is None or not isinstance(value, str) or not value.strip():
                continue

            # Luôn bật wrap cho ô có nội dung để chữ tự xuống dòng.
            alignment = copy.copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment

            width = _merged_width(ws, row, col)
            if width <= 0:
                continue

            font_size = float(cell.font.sz or 11)
            font_factor = max(0.72, min(1.0, 11.0 / font_size))
            chars_per_line = max(5, int(width * 0.95 * font_factor))

            line_count = 0
            for raw_line in (
                str(value).replace("\r\n", "\n").replace("\r", "\n").split("\n")
            ):
                raw_line = raw_line.expandtabs(4)

                if not raw_line:
                    line_count += 1
                    continue

                line_count += max(
                    1,
                    (len(raw_line) + chars_per_line - 1) // chars_per_line,
                )

            line_height = max(13.5, font_size * 1.30)
            cell_height = line_count * line_height + 4
            required_height = max(required_height, cell_height)

        if required_height > min_height:
            ws.row_dimensions[row].height = min(required_height, max_height)
        elif ws.row_dimensions[row].height is None:
            ws.row_dimensions[row].height = min_height


def prepare_product_area(ws, product_count):
    """
    Xử lý vùng bảng trong file mẫu duyet_gia.xlsx hiện tại.

    Cấu trúc mẫu:
        17 = tiêu đề "Hạng mục"
        18 = dòng sản phẩm mẫu
        19..23 = 5 dòng đánh giá
        24 = đoạn đề xuất
        25.. = phần ký duyệt/chữ ký

    openpyxl không tự cập nhật các merged cells khi insert/delete rows,
    vì vậy phải unmerge các vùng phía dưới trước khi chèn dòng rồi
    merge lại với số dòng mới.
    """
    product_template_row = 18
    first_evaluation_row = 19

    if product_count <= 0:
        product_count = 1

    delta = product_count - 1

    if delta <= 0:
        return product_template_row

    # Lưu các merge nằm phía dưới dòng sản phẩm để dịch theo số dòng chèn.
    merge_ranges_to_shift = []
    for merge_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = merge_range.bounds
        if min_row >= first_evaluation_row:
            merge_ranges_to_shift.append((min_col, min_row, max_col, max_row))

    # Phải unmerge trước khi insert, nếu không các merge cũ sẽ đè vào
    # vùng đánh giá/sản phẩm mới.
    for min_col, min_row, max_col, max_row in merge_ranges_to_shift:
        ws.unmerge_cells(
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        )

    # Chèn thêm dòng ngay trước phần đánh giá.
    ws.insert_rows(first_evaluation_row, delta)

    # Copy định dạng của dòng sản phẩm mẫu cho các dòng mới.
    for row in range(first_evaluation_row, first_evaluation_row + delta):
        copy_row_style(
            ws,
            product_template_row,
            row,
            min_col=1,
            max_col=10,
        )

    # Khôi phục các vùng merge ở vị trí mới.
    for min_col, min_row, max_col, max_row in merge_ranges_to_shift:
        shifted_min_row = min_row + delta
        shifted_max_row = max_row + delta

        ws.merge_cells(
            f"{get_column_letter(min_col)}{shifted_min_row}:"
            f"{get_column_letter(max_col)}{shifted_max_row}"
        )

    return product_template_row


def restore_template_drawing(template_path, output_path):
    """
    Khôi phục Shape/Connector (đường gạch ngang) từ file Excel mẫu.
    """

    import xml.etree.ElementTree as ET

    if not os.path.exists(template_path) or not os.path.exists(output_path):
        return

    drawing_name = "xl/drawings/drawing1.xml"
    rels_name = "xl/worksheets/_rels/sheet1.xml.rels"
    sheet_name = "xl/worksheets/sheet1.xml"
    content_types_name = "[Content_Types].xml"

    NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    NS_PKGREL = "http://schemas.openxmlformats.org/package/2006/relationships"

    NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"

    ET.register_namespace("", NS_MAIN)
    ET.register_namespace("r", NS_REL)

    # ==========================================
    # Lấy drawing nguyên bản từ file mẫu
    # ==========================================
    with zipfile.ZipFile(template_path, "r") as template_zip:
        if drawing_name not in template_zip.namelist():
            return

        drawing_data = template_zip.read(drawing_name)

    temp_path = output_path + ".tmp_restore.xlsx"

    # ==========================================
    # Mở file Excel đã xuất
    # ==========================================
    with (
        zipfile.ZipFile(output_path, "r") as src_zip,
        zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as dst_zip,
    ):
        output_names = set(src_zip.namelist())

        # ======================================
        # 1. Relationship của drawing
        # ======================================
        if rels_name in output_names:
            rel_root = ET.fromstring(src_zip.read(rels_name))

        else:
            rel_root = ET.Element(f"{{{NS_PKGREL}}}Relationships")

        drawing_rel = None

        for rel in rel_root.findall(f"{{{NS_PKGREL}}}Relationship"):
            if rel.get("Type", "").endswith("/drawing"):
                drawing_rel = rel
                break

        if drawing_rel is None:
            used_ids = {
                rel.get("Id")
                for rel in rel_root.findall(f"{{{NS_PKGREL}}}Relationship")
            }

            n = 1

            while f"rId{n}" in used_ids:
                n += 1

            drawing_rel = ET.SubElement(
                rel_root,
                f"{{{NS_PKGREL}}}Relationship",
                {
                    "Id": f"rId{n}",
                    "Type": (
                        "http://schemas.openxmlformats.org/"
                        "officeDocument/2006/"
                        "relationships/drawing"
                    ),
                    "Target": "../drawings/drawing1.xml",
                },
            )

        else:
            drawing_rel.set("Target", "../drawings/drawing1.xml")

        drawing_rid = drawing_rel.get("Id")

        rels_data = ET.tostring(
            rel_root,
            encoding="utf-8",
            xml_declaration=True,
        )

        # ======================================
        # 2. Gắn drawing vào sheet
        # ======================================
        sheet_root = ET.fromstring(src_zip.read(sheet_name))

        drawing_element = sheet_root.find(f"{{{NS_MAIN}}}drawing")

        if drawing_element is None:
            drawing_element = ET.Element(
                f"{{{NS_MAIN}}}drawing",
                {f"{{{NS_REL}}}id": drawing_rid},
            )

            children = list(sheet_root)
            insert_index = len(children)

            for i, child in enumerate(children):
                if child.tag in (
                    f"{{{NS_MAIN}}}pageMargins",
                    f"{{{NS_MAIN}}}pageSetup",
                    f"{{{NS_MAIN}}}headerFooter",
                    f"{{{NS_MAIN}}}rowBreaks",
                    f"{{{NS_MAIN}}}colBreaks",
                ):
                    insert_index = i + 1

            sheet_root.insert(insert_index, drawing_element)

        else:
            drawing_element.set(f"{{{NS_REL}}}id", drawing_rid)

        sheet_data = ET.tostring(
            sheet_root,
            encoding="utf-8",
            xml_declaration=True,
        )

        # ======================================
        # 3. Content Type
        # ======================================
        content_types_root = ET.fromstring(src_zip.read(content_types_name))

        has_drawing_content_type = any(
            item.get("PartName") == "/xl/drawings/drawing1.xml"
            for item in content_types_root.findall(f"{{{NS_CT}}}Override")
        )

        if not has_drawing_content_type:
            ET.SubElement(
                content_types_root,
                f"{{{NS_CT}}}Override",
                {
                    "PartName": "/xl/drawings/drawing1.xml",
                    "ContentType": "application/vnd.openxmlformats-officedocument.drawing+xml",
                },
            )

        content_types_data = ET.tostring(
            content_types_root,
            encoding="utf-8",
            xml_declaration=True,
        )

        # ======================================
        # 4. Ghi lại workbook
        # ======================================
        for item in src_zip.infolist():
            name = item.filename

            if name == rels_name:
                dst_zip.writestr(item, rels_data)

            elif name == sheet_name:
                dst_zip.writestr(item, sheet_data)

            elif name == content_types_name:
                dst_zip.writestr(item, content_types_data)

            elif name == drawing_name:
                dst_zip.writestr(item, drawing_data)

            else:
                dst_zip.writestr(item, src_zip.read(name))

        # ======================================
        # QUAN TRỌNG:
        # Nếu file output chưa có drawing1.xml
        # thì thêm nó vào
        # ======================================
        if drawing_name not in output_names:
            drawing_info = zipfile.ZipInfo(drawing_name)

            drawing_info.compress_type = zipfile.ZIP_DEFLATED

            dst_zip.writestr(drawing_info, drawing_data)

    os.replace(temp_path, output_path)


def export_duyet_gia(data, template_path, output_path):
    if not os.path.exists(template_path):
        raise FileNotFoundError(
            f"Không tìm thấy file mẫu:\n{os.path.abspath(template_path)}"
        )

    wb = load_workbook(template_path)

    if "Bộ 01" not in wb.sheetnames:
        raise ValueError("File Excel không có sheet 'Bộ 01'.")

    ws = wb["Bộ 01"]

    vendors = data["vendors"]
    products = data["products"]

    if not vendors:
        raise ValueError("Phải có ít nhất 1 nhà cung cấp.")

    if len(vendors) > MAX_VENDORS:
        raise ValueError(f"Mẫu này chỉ hỗ trợ tối đa {MAX_VENDORS} nhà cung cấp.")

    if not products:
        raise ValueError("Phải có ít nhất 1 hạng mục.")

    # --------------------------------------------------------
    # 1. THÔNG TIN CHUNG
    # --------------------------------------------------------

    so_duyet = data["so_duyet"].strip()
    ngay = data["ngay"].strip()
    thang = data["thang"].strip()
    nam = data["nam"].strip()

    ten_duyet_gia = data["ten_duyet_gia"].strip()
    muc_dich = data["muc_dich_mua_sam"].strip()
    dia_diem = data["dia_diem"].strip()

    ws["A4"] = f"Số: {so_duyet}/KTCN"
    ws["F4"] = f"{dia_diem}, ngày {ngay} tháng {thang} năm {nam}"
    ws["A6"] = f"Ngày {ngay} tháng {thang} năm {nam}"

    ws["F5"] = f"ĐÁNH GIÁ VÀ ĐỀ XUẤT\nLựa chọn nhà cung cấp mua sắm\n{ten_duyet_gia}"

    # --------------------------------------------------------
    # 2. DANH SÁCH NHÀ CUNG CẤP
    # --------------------------------------------------------

    vendor_sentence = (
        f"        Căn cứ yêu cầu mua sắm thiết bị, dịch vụ {muc_dich}, "
        "Phòng KTCN đã yêu cầu báo giá cung cấp của "
        f"{len(vendors):02d} nhà cung cấp sau:"
    )

    for i, vendor in enumerate(vendors, start=1):
        vendor_sentence += f"\n{i}. {vendor['name'].strip()}"

    ws["A13"] = vendor_sentence

    # --------------------------------------------------------
    # 3. TÊN NHÀ CUNG CẤP TRÊN BẢNG GIÁ
    # --------------------------------------------------------
    #
    # File mẫu thực tế đặt {{vendor1}}, {{vendor2}}, {{vendor3}}
    # ở E16:G16, không phải E19:G19.
    for col in range(5, 8):
        cell = ws.cell(16, col)
        if isinstance(cell, MergedCell):
            raise ValueError(  # noqa: TRY004
                f"Ô {cell.coordinate} đang là MergedCell. "
                "Kiểm tra lại merge ở hàng 16 của file mẫu."
            )
        cell.value = None

    for index, vendor in enumerate(vendors):
        col = 5 + index
        ws.cell(16, col).value = vendor["name"].strip()

    # Xóa tên NCC không sử dụng.
    for col in range(5 + len(vendors), 8):
        ws.cell(16, col).value = None

    # --------------------------------------------------------
    # 4. KHU VỰC SẢN PHẨM
    # --------------------------------------------------------

    first_product_row = 18
    prepare_product_area(ws, len(products))

    # Sau khi thêm dòng sản phẩm, phần đánh giá bắt đầu ngay sau
    # danh sách sản phẩm.
    evaluation_start = first_product_row + len(products)

    for index, product in enumerate(products):
        row = first_product_row + index

        # STT
        ws.cell(row, 1).value = f"1.{index + 1}"

        # Tên
        ws.cell(row, 2).value = product["name"].strip()

        # Số lượng
        quantity = product["quantity"].strip()
        try:
            quantity_value = float(quantity)
            if quantity_value.is_integer():
                quantity_value = int(quantity_value)
        except ValueError:
            quantity_value = quantity
        ws.cell(row, 3).value = quantity_value

        # ĐVT
        ws.cell(row, 4).value = product["unit"].strip()

        # Giá các nhà cung cấp
        prices = []
        for vendor_index in range(MAX_VENDORS):
            col = 5 + vendor_index

            if vendor_index < len(vendors):
                price = parse_number(product["prices"][vendor_index])
                ws.cell(row, col).value = price
                ws.cell(row, col).number_format = "#,##0"
                prices.append(price)
            else:
                ws.cell(row, col).value = None

        if not prices:
            raise ValueError(f"Hạng mục {index + 1} chưa có giá.")

        # Giá thấp nhất
        min_price = min(prices)
        ws.cell(row, 8).value = min_price
        ws.cell(row, 8).number_format = "#,##0"

        # NCC có giá thấp nhất cho riêng hạng mục
        min_index = prices.index(min_price)
        ws.cell(row, 9).value = vendors[min_index]["name"].strip()

        ws.cell(row, 10).value = product.get("note", "").strip()

    # --------------------------------------------------------
    # 5. ĐÁNH GIÁ NCC
    # --------------------------------------------------------

    row_capacity = evaluation_start
    row_technical = evaluation_start + 1
    row_delivery = evaluation_start + 2
    row_time = evaluation_start + 3
    row_payment = evaluation_start + 4
    row_proposal = evaluation_start + 5

    labels = {
        row_capacity: "Đánh giá về năng lực và KN",
        row_technical: "Đánh giá về đáp ứng yêu cầu KT",
        row_delivery: "Địa điểm giao hàng/thực hiện",
        row_time: "Thời gian thực hiện hợp đồng",
        row_payment: "Thời hạn thanh toán",
    }

    for row, label in labels.items():
        ws.cell(row, 1).value = row - evaluation_start + 2
        ws.cell(row, 2).value = label

    for vendor_index, vendor in enumerate(vendors):
        col = 5 + vendor_index

        ws.cell(row_capacity, col).value = vendor["nang_luc"].strip()
        ws.cell(row_technical, col).value = vendor["ky_thuat"].strip()
        ws.cell(row_delivery, col).value = vendor["dia_diem"].strip()
        ws.cell(row_time, col).value = vendor["thoi_gian"].strip()
        ws.cell(row_payment, col).value = vendor["cach_thanh_toan"].strip()

    # Xóa cột E-G cho NCC không sử dụng.
    for vendor_index in range(len(vendors), MAX_VENDORS):
        col = 5 + vendor_index
        for row in labels:
            ws.cell(row, col).value = None

    # --------------------------------------------------------
    # 6. CHỌN NCC TỔNG THỂ
    # --------------------------------------------------------

    vendor_totals = [0] * len(vendors)

    for product in products:
        for i in range(len(vendors)):
            vendor_totals[i] += parse_number(product["prices"][i])

    selected_index = vendor_totals.index(min(vendor_totals))
    selected_vendor = vendors[selected_index]["name"].strip()
    selected_total = vendor_totals[selected_index]

    # --------------------------------------------------------
    # 7. ĐOẠN ĐỀ XUẤT
    # --------------------------------------------------------

    proposal = (
        "        Phòng KTCN đề xuất Thủ trưởng Nhà máy cho mua các hạng mục "
        f"từ {selected_vendor} vì có tổng giá chào thấp nhất "
        "trong các nhà cung cấp được đánh giá, đáp ứng các yêu cầu "
        "về vật tư, kỹ thuật và khả năng thực hiện."
    )

    # Dòng đề xuất trong mẫu là merge A:J.
    target_merge = f"A{row_proposal}:J{row_proposal}"
    already_merged = any(str(rng) == target_merge for rng in ws.merged_cells.ranges)

    if not already_merged:
        # Trường hợp file mẫu đã bị chỉnh sửa trước đó.
        for rng in list(ws.merged_cells.ranges):
            min_col, min_row, max_col, max_row = rng.bounds
            if (
                min_row == row_proposal
                and max_row == row_proposal
                and min_col <= 1 <= max_col
            ):
                ws.unmerge_cells(str(rng))
        ws.merge_cells(target_merge)

    ws.cell(row_proposal, 1).value = proposal

    # --------------------------------------------------------
    # 8. GHI TỔNG GIÁ
    # --------------------------------------------------------

    from openpyxl.comments import Comment

    for i, total in enumerate(vendor_totals):
        col = 5 + i
        ws.cell(16, col).comment = Comment(
            f"Tổng giá chưa VAT: {format_money(total)} đồng",
            "Python",
        )

    # --------------------------------------------------------
    # 9. FORMAT
    # --------------------------------------------------------

    # Tự động xuống dòng + tự động tăng chiều cao hàng theo nội dung.
    # Có hỗ trợ cả ô Merge vì openpyxl không AutoFit Row Height chuẩn.
    auto_fit_row_heights(
        ws,
        min_row=17,
        max_row=ws.max_row,
        min_height=15.0,
        max_height=409.0,
    )

    ws.sheet_view.showGridLines = False

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)

    # Khôi phục các đường gạch ngang/Shape từ file mẫu
    restore_template_drawing(template_path, output_path)

    return {
        "output_path": os.path.abspath(output_path),
        "selected_vendor": selected_vendor,
        "selected_total": selected_total,
        "vendor_totals": vendor_totals,
    }


# ============================================================
# GIAO DIỆN TKINTER
# ============================================================


class ScrollableFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)

        self.canvas = tk.Canvas(self, highlightthickness=0)

        self.scrollbar = ttk.Scrollbar(
            self, orient="vertical", command=self.canvas.yview
        )

        self.inner = ttk.Frame(self.canvas)

        self.inner.bind(
            "<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")),
        )

        self.window_id = self.canvas.create_window(
            (0, 0), window=self.inner, anchor="nw"
        )

        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")

        self.canvas.bind("<Configure>", self._resize_inner)

        self.canvas.bind_all("<MouseWheel>", self._on_mousewheel)

    def _resize_inner(self, event):
        self.canvas.itemconfig(self.window_id, width=event.width)

    def _on_mousewheel(self, event):
        self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")


class DuyetGiaApp(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title("Lập phiếu phê duyệt đánh giá và đề xuất")
        self.geometry("1250x850")
        self.minsize(1000, 700)

        self.vendor_vars = []
        self.product_rows = []

        self.create_style()
        self.create_ui()

        self.add_default_vendors()
        self.add_product()

    # --------------------------------------------------------
    # STYLE
    # --------------------------------------------------------

    def create_style(self):
        style = ttk.Style(self)

        try:
            style.theme_use("clam")
        except Exception:  # noqa: BLE001, S110
            pass

        style.configure("Title.TLabel", font=("Arial", 16, "bold"))

        style.configure("Section.TLabel", font=("Arial", 12, "bold"))

        style.configure("Header.TLabel", font=("Arial", 10, "bold"))

        style.configure("Action.TButton", font=("Arial", 10, "bold"), padding=8)

    # --------------------------------------------------------
    # UI
    # --------------------------------------------------------

    def create_ui(self):

        header = ttk.Frame(self, padding=10)
        header.pack(fill="x")

        ttk.Label(
            header, text="PHÊ DUYỆT ĐÁNH GIÁ VÀ ĐỀ XUẤT", style="Title.TLabel"
        ).pack(side="left")

        ttk.Button(header, text="📂 Chọn file mẫu", command=self.choose_template).pack(
            side="right"
        )

        self.template_label = ttk.Label(
            header, text=f"Mẫu: {os.path.abspath(TEMPLATE_FILE)}"
        )

        ttk.Button(
            header,
            text=f"📂 Chọn file hợp đồng mẫu",
            command=self.choose_contract_template,
        ).pack(side="right")

        self.template_label = ttk.Label(
            header, text=f"Mẫu: {os.path.abspath(TEMPLATE_FILE)}"
        )

        self.template_label.pack(side="right", padx=20)

        self.main = ScrollableFrame(self)
        self.main.pack(fill="both", expand=True, padx=10)

        self.create_general_section()
        self.create_vendor_section()
        self.create_product_section()
        self.create_action_section()

    # --------------------------------------------------------
    # THÔNG TIN CHUNG
    # --------------------------------------------------------

    def create_general_section(self):

        frame = ttk.LabelFrame(self.main.inner, text=" 1. Thông tin chung ", padding=10)
        frame.pack(fill="x", pady=(0, 10))

        for col in range(6):
            frame.columnconfigure(col, weight=1)

        ttk.Label(frame, text="Số duyệt:", style="Header.TLabel").grid(
            row=0, column=0, sticky="w", padx=5, pady=5
        )

        self.so_duyet = ttk.Entry(frame)
        self.so_duyet.grid(row=0, column=1, sticky="ew", padx=5, pady=5)

        ttk.Label(frame, text="Ngày:", style="Header.TLabel").grid(
            row=0, column=2, sticky="w", padx=5, pady=5
        )

        self.ngay = ttk.Entry(frame, width=8)
        self.ngay.grid(row=0, column=3, sticky="ew", padx=5, pady=5)

        ttk.Label(frame, text="Tháng:", style="Header.TLabel").grid(
            row=0, column=4, sticky="w", padx=5, pady=5
        )

        self.thang = ttk.Entry(frame, width=8)
        self.thang.grid(row=0, column=5, sticky="ew", padx=5, pady=5)

        ttk.Label(frame, text="Năm:", style="Header.TLabel").grid(
            row=1, column=0, sticky="w", padx=5, pady=5
        )

        self.nam = ttk.Entry(frame)
        self.nam.grid(row=1, column=1, sticky="ew", padx=5, pady=5)

        now = datetime.now()  # noqa: DTZ005

        self.ngay.insert(0, f"{now.day:02d}")
        self.thang.insert(0, f"{now.month:02d}")
        self.nam.insert(0, str(now.year))

        ttk.Label(frame, text="Địa điểm:", style="Header.TLabel").grid(
            row=1, column=2, sticky="w", padx=5, pady=5
        )

        self.dia_diem = ttk.Entry(frame)
        self.dia_diem.insert(0, "Lào Cai")
        self.dia_diem.grid(row=1, column=3, columnspan=3, sticky="ew", padx=5, pady=5)

        ttk.Label(frame, text="Tên nội dung duyệt giá:", style="Header.TLabel").grid(
            row=2, column=0, sticky="nw", padx=5, pady=5
        )

        self.ten_duyet_gia = tk.Text(frame, height=3, wrap="word")
        self.ten_duyet_gia.grid(
            row=2, column=1, columnspan=5, sticky="ew", padx=5, pady=5
        )

        ttk.Label(
            frame, text="Mục đích / yêu cầu mua sắm:", style="Header.TLabel"
        ).grid(row=3, column=0, sticky="nw", padx=5, pady=5)

        self.muc_dich = tk.Text(frame, height=3, wrap="word")
        self.muc_dich.grid(row=3, column=1, columnspan=5, sticky="ew", padx=5, pady=5)

        self.loai_hd = tk.StringVar(value="Dịch vụ")

        ttk.Radiobutton(
            frame,
            text="Dịch vụ",
            variable=self.loai_hd,
            value="Dịch vụ",
            command=self.change_contract_type,
        ).grid(row=5, column=0, columnspan=2, sticky="w", padx=20, pady=5)

        ttk.Radiobutton(
            frame,
            text="Mua bán",
            variable=self.loai_hd,
            value="Mua bán",
            command=self.change_contract_type,
        ).grid(row=5, column=2, columnspan=2, sticky="w", padx=20, pady=5)

        ttk.Label(frame, text="Mã Hợp Đồng:", style="Header.TLabel").grid(
            row=4, column=0, sticky="w", padx=5, pady=5
        )

        self.ma_hd = ttk.Entry(frame)
        self.ma_hd.grid(row=4, column=1, sticky="ew", padx=5, pady=5)

    def change_contract_type(self):
        if self.loai_hd.get() == "Dịch vụ":
            self.sale_frame.pack_forget()

            self.service_frame.pack(
                fill="both",
                expand=True,
                padx=0,
                pady=0,
            )

            self.service_info_frame.grid(
                row=2,
                column=0,
                columnspan=4,
                sticky="ew",
            )

            self.template_file = DEFAULT_SERVICE_TEMPLATE

        else:
            self.service_frame.pack_forget()

            self.sale_frame.pack(
                fill="both",
                expand=True,
                padx=0,
                pady=0,
            )

            self.service_info_frame.grid_remove()

            self.template_file = DEFAULT_SALE_TEMPLATE

        self.template_label.config(text=self.template_file)

    # --------------------------------------------------------
    # NHÀ CUNG CẤP
    # --------------------------------------------------------

    def create_vendor_section(self):

        frame = ttk.LabelFrame(self.main.inner, text=" 2. Nhà cung cấp ", padding=10)
        frame.pack(fill="x", pady=(0, 10))

        headers = [
            "STT",
            "Tên nhà cung cấp",
            "Năng lực",
            "Đáp ứng kỹ thuật",
            "Địa điểm giao hàng",
            "Thời gian thực hiện",
            "Thời hạn thanh toán",
        ]

        for col, text in enumerate(headers):
            ttk.Label(frame, text=text, style="Header.TLabel").grid(
                row=0, column=col, sticky="nsew", padx=3, pady=3
            )

        for col in range(len(headers)):
            frame.columnconfigure(col, weight=1)

        self.vendor_container = frame

    def add_default_vendors(self):

        defaults = [
            {
                "name": "Nhà cung cấp 1",
                "nang_luc": "Đạt",
                "ky_thuat": "Đạt",
                "dia_diem": "Lào Cai",
                "thoi_gian": "Trong vòng 04 tuần kể từ ngày Bên A tạm ứng.",
                "cach_thanh_toan": "Theo thỏa thuận trong hợp đồng.",
            },
            {
                "name": "Nhà cung cấp 2",
                "nang_luc": "Đạt",
                "ky_thuat": "Đạt",
                "dia_diem": "Lào Cai",
                "thoi_gian": "Trong vòng 06 tuần kể từ ngày Bên A tạm ứng.",
                "cach_thanh_toan": "Theo thỏa thuận trong hợp đồng.",
            },
            {
                "name": "Nhà cung cấp 3",
                "nang_luc": "Đạt",
                "ky_thuat": "Đạt",
                "dia_diem": "Lào Cai",
                "thoi_gian": "Trong vòng 06 tuần kể từ ngày Bên A tạm ứng.",
                "cach_thanh_toan": "Theo thỏa thuận trong hợp đồng.",
            },
        ]

        for item in defaults:
            self.add_vendor(item)

    def add_vendor(self, data=None):

        if len(self.vendor_vars) >= MAX_VENDORS:
            messagebox.showinfo(
                "Thông báo",
                f"Mẫu Excel hiện tại hỗ trợ tối đa {MAX_VENDORS} nhà cung cấp.",
            )
            return

        data = data or {}

        row = len(self.vendor_vars) + 1

        variables = {
            "name": tk.StringVar(value=data.get("name", "")),
            "nang_luc": tk.StringVar(value=data.get("nang_luc", "Đạt")),
            "ky_thuat": tk.StringVar(value=data.get("ky_thuat", "Đạt")),
            "dia_diem": tk.StringVar(value=data.get("dia_diem", "")),
            "thoi_gian": tk.StringVar(value=data.get("thoi_gian", "")),
            "cach_thanh_toan": tk.StringVar(value=data.get("cach_thanh_toan", "")),
        }

        ttk.Label(self.vendor_container, text=str(row)).grid(
            row=row, column=0, padx=3, pady=3
        )

        ttk.Entry(self.vendor_container, textvariable=variables["name"]).grid(
            row=row, column=1, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.vendor_container, textvariable=variables["nang_luc"]).grid(
            row=row, column=2, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.vendor_container, textvariable=variables["ky_thuat"]).grid(
            row=row, column=3, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.vendor_container, textvariable=variables["dia_diem"]).grid(
            row=row, column=4, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.vendor_container, textvariable=variables["thoi_gian"]).grid(
            row=row, sticky="ew", column=5, padx=3, pady=3
        )

        ttk.Entry(
            self.vendor_container, textvariable=variables["cach_thanh_toan"]
        ).grid(row=row, column=6, sticky="ew", padx=3, pady=3)

        self.vendor_vars.append(variables)

    # --------------------------------------------------------
    # SẢN PHẨM
    # --------------------------------------------------------

    def create_product_section(self):

        frame = ttk.LabelFrame(
            self.main.inner, text=" 3. Hạng mục / vật tư ", padding=10
        )
        frame.pack(fill="x", pady=(0, 10))

        headers = [
            "STT",
            "Tên hạng mục",
            "Số lượng",
            "ĐVT",
            "Giá NCC 1",
            "Giá NCC 2",
            "Giá NCC 3",
            "Ghi chú",
            "",
        ]

        for col, text in enumerate(headers):
            ttk.Label(frame, text=text, style="Header.TLabel").grid(
                row=0, column=col, sticky="nsew", padx=3, pady=3
            )

        widths = [5, 28, 10, 8, 15, 15, 15, 20, 8]

        for col, weight in enumerate(widths):
            frame.columnconfigure(col, weight=weight)

        self.product_container = frame

        ttk.Button(frame, text="➕ Thêm hạng mục", command=self.add_product).grid(
            row=999, column=0, columnspan=9, sticky="w", padx=3, pady=8
        )

    def add_product(self):

        row = len(self.product_rows) + 1

        vars_ = {
            "name": tk.StringVar(),
            "quantity": tk.StringVar(value="1"),
            "unit": tk.StringVar(value="Cái"),
            "price1": tk.StringVar(),
            "price2": tk.StringVar(),
            "price3": tk.StringVar(),
            "note": tk.StringVar(),
        }

        ttk.Label(self.product_container, text=str(row)).grid(
            row=row, column=0, padx=3, pady=3
        )

        ttk.Entry(self.product_container, textvariable=vars_["name"]).grid(
            row=row, column=1, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.product_container, textvariable=vars_["quantity"]).grid(
            row=row, column=2, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.product_container, textvariable=vars_["unit"]).grid(
            row=row, column=3, sticky="ew", padx=3, pady=3
        )

        for col, key in [
            (4, "price1"),
            (5, "price2"),
            (6, "price3"),
        ]:
            ttk.Entry(self.product_container, textvariable=vars_[key]).grid(
                row=row, column=col, sticky="ew", padx=3, pady=3
            )

        ttk.Entry(self.product_container, textvariable=vars_["note"]).grid(
            row=row, column=7, sticky="ew", padx=3, pady=3
        )

        if row > 1:
            ttk.Button(
                self.product_container,
                text="Xóa",
                command=lambda r=row: self.remove_product(r),
            ).grid(row=row, column=8, padx=3, pady=3)

        self.product_rows.append(vars_)

        self.update_add_button()

    def remove_product(self, row_number):

        if len(self.product_rows) <= 1:
            messagebox.showinfo("Thông báo", "Phải có ít nhất 1 hạng mục.")
            return

        # Xóa widget của dòng.
        for widget in self.product_container.grid_slaves(row=row_number):
            widget.destroy()

        # Lưu lại các biến, bỏ phần tử tương ứng.
        self.product_rows.pop(row_number - 1)

        # Vẽ lại toàn bộ phần sản phẩm để STT không bị lệch.
        for widget in self.product_container.grid_slaves():
            info = widget.grid_info()
            try:
                r = int(info["row"])
            except Exception:
                continue

            if r != 0 and r != 999:
                widget.destroy()

        old_rows = self.product_rows
        self.product_rows = []

        for vars_ in old_rows:
            self._render_product_row(vars_)

        self.update_add_button()

    def _render_product_row(self, vars_):

        row = len(self.product_rows) + 1

        ttk.Label(self.product_container, text=str(row)).grid(
            row=row, column=0, padx=3, pady=3
        )

        ttk.Entry(self.product_container, textvariable=vars_["name"]).grid(
            row=row, column=1, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.product_container, textvariable=vars_["quantity"]).grid(
            row=row, column=2, sticky="ew", padx=3, pady=3
        )

        ttk.Entry(self.product_container, textvariable=vars_["unit"]).grid(
            row=row, column=3, sticky="ew", padx=3, pady=3
        )

        for col, key in [
            (4, "price1"),
            (5, "price2"),
            (6, "price3"),
        ]:
            ttk.Entry(self.product_container, textvariable=vars_[key]).grid(
                row=row, column=col, sticky="ew", padx=3, pady=3
            )

        ttk.Entry(self.product_container, textvariable=vars_["note"]).grid(
            row=row, column=7, sticky="ew", padx=3, pady=3
        )

        if row > 1:
            ttk.Button(
                self.product_container,
                text="Xóa",
                command=lambda r=row: self.remove_product(r),
            ).grid(row=row, column=8, padx=3, pady=3)

        self.product_rows.append(vars_)

    def update_add_button(self):
        # Nút thêm được tạo ở row 999, không cần xử lý thêm.
        pass

    # --------------------------------------------------------
    # ACTION
    # --------------------------------------------------------

    def create_action_section(self):

        frame = ttk.Frame(self.main.inner, padding=10)
        frame.pack(fill="x", pady=10)

        ttk.Button(
            frame, text="💾 XUẤT EXCEL", style="Action.TButton", command=self.export
        ).pack(side="left", padx=5)

        ttk.Button(
            frame,
            text="🔍 Xem dữ liệu",
            style="Action.TButton",
            command=self.preview_data,
        ).pack(side="left", padx=5)

        ttk.Button(
            frame,
            text="📁 Mở thư mục output",
            style="Action.TButton",
            command=self.open_output_dir,
        ).pack(side="left", padx=5)

    # --------------------------------------------------------
    # DATA
    # --------------------------------------------------------

    def get_data(self):

        vendors = []

        for vendor in self.vendor_vars:
            name = vendor["name"].get().strip()

            if not name:
                continue

            vendors.append(
                {
                    "name": name,
                    "nang_luc": vendor["nang_luc"].get(),
                    "ky_thuat": vendor["ky_thuat"].get(),
                    "dia_diem": vendor["dia_diem"].get(),
                    "thoi_gian": vendor["thoi_gian"].get(),
                    "cach_thanh_toan": vendor["cach_thanh_toan"].get(),
                }
            )

        products = []

        for index, row in enumerate(self.product_rows, start=1):
            name = row["name"].get().strip()

            if not name:
                continue

            products.append(
                {
                    "id": index,
                    "name": name,
                    "quantity": row["quantity"].get().strip(),
                    "unit": row["unit"].get().strip(),
                    "prices": [
                        row["price1"].get().strip(),
                        row["price2"].get().strip(),
                        row["price3"].get().strip(),
                    ],
                    "note": row["note"].get().strip(),
                }
            )

        # Kiểm tra giá theo số NCC.
        for product in products:
            for i in range(len(vendors)):
                parse_number(product["prices"][i])

        return {
            "so_duyet": self.so_duyet.get(),
            "so_duyet": self.so_duyet.get(),
            "ngay": self.ngay.get(),
            "thang": self.thang.get(),
            "nam": self.nam.get(),
            "ten_duyet_gia": self.ten_duyet_gia.get("1.0", "end").strip(),
            "muc_dich_mua_sam": self.muc_dich.get("1.0", "end").strip(),
            "dia_diem": self.dia_diem.get(),
            "vendors": vendors,
            "products": products,
        }

    # --------------------------------------------------------
    # EXPORT
    # --------------------------------------------------------

    def export(self):

        try:
            data = self.get_data()

            if len(data["vendors"]) == 0:
                raise ValueError("Chưa nhập nhà cung cấp.")

            if len(data["vendors"]) > MAX_VENDORS:
                raise ValueError(f"Tối đa {MAX_VENDORS} nhà cung cấp.")

            if len(data["products"]) == 0:
                raise ValueError("Chưa nhập hạng mục.")

            # Mỗi sản phẩm cần đủ giá cho các NCC.
            for index, product in enumerate(data["products"], start=1):
                for i in range(len(data["vendors"])):
                    if not product["prices"][i].strip():
                        raise ValueError(
                            f"Hạng mục {index} chưa nhập giá của NCC {i + 1}."
                        )

            os.makedirs(DEFAULT_OUTPUT_DIR, exist_ok=True)

            default_name = f"duyet_gia_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"  # noqa: DTZ005

            output_path = filedialog.asksaveasfilename(
                title="Lưu file Excel",
                initialdir=os.path.abspath(DEFAULT_OUTPUT_DIR),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel Workbook", "*.xlsx")],
            )

            if not output_path:
                return

            result = export_duyet_gia(
                data=data, template_path=TEMPLATE_FILE, output_path=output_path
            )

            messagebox.showinfo(
                "Xuất Excel thành công",
                "Đã tạo file:\n\n"
                f"{result['output_path']}\n\n"
                f"Nhà cung cấp được chọn theo tổng giá:\n"
                f"{result['selected_vendor']}\n\n"
                f"Tổng giá chưa VAT:\n"
                f"{format_money(result['selected_total'])} đồng",
            )

            open_file(output_path)

        except Exception as exc:
            messagebox.showerror("Lỗi", str(exc))

    # --------------------------------------------------------
    # PREVIEW
    # --------------------------------------------------------

    def preview_data(self):

        try:
            data = self.get_data()

            text = []

            text.append("===== NHÀ CUNG CẤP =====")

            for i, vendor in enumerate(data["vendors"], start=1):
                text.append(f"{i}. {vendor['name']}")

            text.append("")
            text.append("===== HẠNG MỤC =====")

            totals = [0] * len(data["vendors"])

            for i, product in enumerate(data["products"], start=1):
                text.append(
                    f"{i}. {product['name']} - {product['quantity']} {product['unit']}"
                )

                for vendor_index in range(len(data["vendors"])):
                    price = parse_number(product["prices"][vendor_index])
                    totals[vendor_index] += price

                    text.append(f"   NCC {vendor_index + 1}: {format_money(price)}")

            if totals:
                selected_index = totals.index(min(totals))

                text.append("")
                text.append("===== TỔNG =====")

                for i, total in enumerate(totals, start=1):
                    text.append(f"NCC {i}: {format_money(total)} đồng")

                text.append("")
                text.append(
                    f"Đề xuất theo tổng giá: {data['vendors'][selected_index]['name']}"
                )

            messagebox.showinfo("Xem dữ liệu", "\n".join(text))

        except Exception as exc:
            messagebox.showerror("Lỗi dữ liệu", str(exc))

    # --------------------------------------------------------
    # TEMPLATE
    # --------------------------------------------------------

    def choose_template(self):
        global TEMPLATE_FILE
        path = filedialog.askopenfilename(
            title="Chọn file Excel mẫu", filetypes=[("Excel Workbook", "*.xlsx")]
        )
        if not path:
            return
        TEMPLATE_FILE = path
        self.template_label.config(text=f"Mẫu: {os.path.abspath(TEMPLATE_FILE)}")

    def choose_contract_template(self):
        global TEMPLATE_CONTRACT_FILE
        file_path = filedialog.askopenfilename(
            title="Chọn file Word template",
            filetypes=[
                ("Word files", "*.docx"),
                ("All files", "*.*"),
            ],
        )
        if not file_path:
            return
        self.template_contract_file = file_path
        self.template_contract_file.config(text=file_path)
        TEMPLATE_CONTRACT_FILE = file_path
        self.template_label.config(
            text=f"Mẫu: {os.path.abspath(TEMPLATE_CONTRACT_FILE)}"
        )

    # --------------------------------------------------------
    # OUTPUT
    # --------------------------------------------------------

    def open_output_dir(self):

        path = os.path.abspath(DEFAULT_OUTPUT_DIR)
        os.makedirs(path, exist_ok=True)
        open_file(path)


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    app = DuyetGiaApp()
    app.mainloop()
